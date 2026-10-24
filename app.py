import streamlit as st

import sqlite3

import psycopg
from psycopg.rows import dict_row

import pandas as pd

import hashlib

from pathlib import Path

from datetime import datetime

from io import BytesIO

import base64

from html import escape

import json
import os
import re
import secrets



from openpyxl import load_workbook

from openpyxl.styles import Font, PatternFill, Alignment



# =========================================================

# 1. CONFIGURATION

# =========================================================

st.set_page_config(page_title="Don Bosco Kakuma MEAL Dashboard", page_icon="📊", layout="wide")

APP_DIR = Path(__file__).parent

DB_FILE = APP_DIR / "meal.db"

try:
          DATABASE_URL = st.secrets.get("DATABASE_URL") or os.environ.get("DATABASE_URL")
except FileNotFoundError:
          DATABASE_URL = os.environ.get("DATABASE_URL")

if DATABASE_URL and "sslmode=" not in DATABASE_URL:
          DATABASE_URL = f"{DATABASE_URL}{'&' if '?' in DATABASE_URL else '?'}sslmode=require"

LOGO_FILE = APP_DIR / "ans.png"

REPORT_TEMPLATE_NAME = "END SEPTEMBER 2026 STATISTICAL REPORT.xlsx"
DEFAULT_REPORT_TEMPLATE = APP_DIR.parent.parent / REPORT_TEMPLATE_NAME
PROJECT_REPORT_TEMPLATE = APP_DIR / REPORT_TEMPLATE_NAME
REPORT_TEMPLATE = Path(
          os.environ.get(
                    "REPORT_TEMPLATE_PATH",
                    str(PROJECT_REPORT_TEMPLATE if PROJECT_REPORT_TEMPLATE.exists() else DEFAULT_REPORT_TEMPLATE),
          )
)

CENTRES = ["DBK", "DB1", "DB2", "DB3", "DB4", "DB Kalobeyei ICT", "DB Kalobeyei Buss"]

COURSES = [

          "Computer Basic", "English Basic", "English Intermediate", "ICT II", "ICT III",

          "Fashion", "Dress Making II", "Dress Making III", "Masonry I", "Masonry II",

          "Masonry III", "Plumbing I", "Plumbing III", "Tailoring II", "Tailoring III",

          "Agriculture", "Carpentry III", "Electrical II", "Electrical III",

          "Secretarial Studies", "Other"

]

GENDERS = ["Male", "Female"]

COMPLETION_MONTHS = ["April", "June", "August", "December"]
DROPOUT_QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
# Exact student-facing status labels. Ongoing remains available for learners
# who have not completed or dropped out yet.
STATUS_DETAILS = [
          "Completed in April", "Completed in June", "Completed in August", "Completed in December",
          "Dropped out Q1", "Dropped out Q2", "Dropped out Q3", "Dropped out Q4"
]
STATUSES = ["Ongoing"] + STATUS_DETAILS

NATIONALITIES = ["South Sudanese", "Sudanese", "Somali", "Ethiopian", "Congolese", "Burundian", "Kenyan", "Other", "Not specified"]
DB_ERRORS = (sqlite3.Error, psycopg.Error)
DB_INTEGRITY_ERRORS = (sqlite3.IntegrityError, psycopg.IntegrityError)
LEARNER_UPDATE_ERRORS = DB_ERRORS + (PermissionError,)
REPORT_NATIONALITY_ALIASES = {
          "Sud": {"sudanese", "sudan"},
          "S.Sud": {"south sudanese", "south sudan", "s sudanese", "s sud", "south sadanese"},
          "Som": {"somali", "somalia"},
          "Ethi": {"ethiopian", "ethiopia"},
          "Cong": {"congolese", "congo", "congolese drc"},
          "Bur": {"burundian", "burundi"},
          "Kenya Others": {"kenya other", "kenya others", "kenyan other", "kenyan others"},
          "Host": {"kenya", "kenyan", "host"},
          "Uga": {"ugandan", "uganda"},
          "Yemen": {"yemen", "yemeni"},
          "Rwan": {"rwandese", "rwandan", "rwanda"},
          "Afg": {"afghan", "afghani", "afghanistan"},
          "Mag": {"madagasy", "malagasy", "madagascar"},
          "Eritr": {"eritrean", "eritrea"},
          "Other / Not specified": {"", "other", "not specified", "unknown", "n/a"},
}



# =========================================================

# 2. STYLING

# =========================================================

st.markdown("""

<style>

.stApp{background:#F4F7FB;color:#203047}.block-container{padding-top:1.3rem;padding-bottom:2rem;max-width:1500px}

h1,h2,h3{color:#12345A!important;font-weight:700!important}
p,li,label{color:#35475D}
[data-testid="stCaptionContainer"]{color:#708197}
[data-testid="stForm"]{background:linear-gradient(145deg,#fff,#F9FBFE);border:1px solid #DCE4EF;border-radius:16px;padding:1.2rem 1.3rem;box-shadow:0 10px 30px rgba(18,52,90,.055)}
[data-testid="stTextInput"] input,[data-testid="stNumberInput"] input,[data-testid="stSelectbox"]>div,[data-testid="stMultiSelect"]>div,[data-testid="stDateInput"] input{border-radius:9px}
[data-testid="stTabs"] [data-baseweb="tab-list"]{gap:.35rem;border-bottom:1px solid #DCE4EF}
[data-testid="stTabs"] [data-baseweb="tab"]{height:44px;border-radius:9px 9px 0 0;color:#52657B;font-weight:600}
[data-testid="stTabs"] [aria-selected="true"]{color:#12345A!important}
[data-testid="stDataFrame"],[data-testid="stTable"],[data-testid="stPlotlyChart"],[data-testid="stVegaLiteChart"]{border:1px solid #DCE4EF;border-radius:13px;overflow:hidden;background:#fff}
[data-testid="stExpander"]{background:#fff;border:1px solid #DCE4EF;border-radius:12px}
.welcome-hero{position:relative;overflow:hidden;min-height:196px;display:flex;align-items:center;justify-content:space-between;gap:1.5rem;padding:2rem 2.2rem;border-radius:22px;background:radial-gradient(circle at 83% 30%,rgba(69,145,194,.28),transparent 28%),linear-gradient(118deg,#0B1E34 0%,#12345A 54%,#1D527D 100%);border:1px solid rgba(212,160,23,.4);border-bottom:5px solid #D4A017;box-shadow:0 18px 40px rgba(18,52,90,.16);margin:.15rem 0 1.25rem}
.welcome-kicker{display:inline-flex;align-items:center;gap:.5rem;color:#F4D16A;font-size:.69rem;font-weight:800;letter-spacing:.16em;text-transform:uppercase}
.welcome-status{width:8px;height:8px;border-radius:50%;background:#60D6A5;box-shadow:0 0 0 4px rgba(96,214,165,.14)}
.welcome-title{color:#fff;font-size:clamp(1.8rem,3.2vw,2.8rem);font-weight:800;letter-spacing:-.035em;line-height:1.08;margin:.8rem 0 .55rem}
.welcome-copy{max-width:660px;color:#C5D5E6;font-size:.96rem;line-height:1.65}
.welcome-chart{height:112px;min-width:190px;display:flex;align-items:flex-end;justify-content:center;gap:10px;padding:1rem 1.2rem;border:1px solid rgba(255,255,255,.14);border-radius:16px;background:rgba(255,255,255,.06)}
.welcome-chart span{width:18px;border-radius:6px 6px 2px 2px;background:linear-gradient(180deg,#F4D16A,#E9B93D);box-shadow:0 4px 14px rgba(212,160,23,.17)}
.welcome-chart span:nth-child(1){height:28%}.welcome-chart span:nth-child(2){height:49%}.welcome-chart span:nth-child(3){height:38%}.welcome-chart span:nth-child(4){height:68%}.welcome-chart span:nth-child(5){height:55%}.welcome-chart span:nth-child(6){height:86%}
.welcome-section-title{color:#12345A;font-size:1.02rem;font-weight:750;margin:.6rem 0 .15rem}
.welcome-section-copy{color:#718197;font-size:.84rem;margin-bottom:.75rem}
@media(max-width:720px){.welcome-hero{min-height:0;padding:1.5rem}.welcome-chart{display:none}}

[data-testid="stSidebar"]{background:linear-gradient(180deg,#0B1E34 0%,#12345A 52%,#102947 100%);border-right:1px solid rgba(212,160,23,.3)}
[data-testid="stSidebar"] [data-testid="stSidebarContent"]{padding:1.1rem .9rem}
[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p{color:#DCE7F3}
[data-testid="stSidebar"] [data-testid="stImage"]{display:flex;justify-content:center;margin:.2rem 0 .9rem}
[data-testid="stSidebar"] .stButton{margin:.12rem 0}
[data-testid="stSidebar"] .stButton button{width:100%;min-height:42px;border-radius:10px;text-align:left;padding:.55rem .8rem;font-weight:600;transition:background .16s ease,border-color .16s ease,transform .16s ease}
[data-testid="stSidebar"] .stButton button[kind="secondary"]{background:rgba(255,255,255,.045);color:#E7EEF7;border:1px solid rgba(220,231,243,.1)}
[data-testid="stSidebar"] .stButton button[kind="secondary"]:hover{background:rgba(255,255,255,.12);border-color:rgba(244,209,106,.45);color:#fff;transform:translateX(2px)}
[data-testid="stSidebar"] .stButton button[kind="primary"]{background:linear-gradient(105deg,#F4D16A,#E9B93D);color:#102947;border:1px solid #F4D16A;box-shadow:0 5px 16px rgba(212,160,23,.2)}
[data-testid="stSidebar"] .stButton button[kind="primary"]:hover{background:#FFE28A;border-color:#FFE28A;color:#102947}
[data-testid="stSidebar"] [data-testid="stVerticalBlock"]{gap:.45rem}
[data-testid="stSidebar"] hr{border-color:rgba(220,231,243,.15);margin:.9rem 0}
[data-testid="stSidebar"] [data-testid="stCaptionContainer"]{color:#8FA7C0}
.meal-brand{padding:.45rem .35rem .9rem;text-align:center}
.meal-brand-kicker{color:#F4D16A;font-size:.66rem;font-weight:800;letter-spacing:.19em;text-transform:uppercase}
.meal-brand-title{color:#fff;font-size:1.2rem;font-weight:800;letter-spacing:.02em;margin-top:.14rem}
.meal-brand-subtitle{color:#AFC2D7;font-size:.72rem;letter-spacing:.12em;margin-top:.1rem;text-transform:uppercase}
.meal-account{background:linear-gradient(135deg,rgba(255,255,255,.11),rgba(255,255,255,.045));border:1px solid rgba(220,231,243,.13);border-radius:13px;padding:.78rem .85rem;margin:.35rem 0 1rem}
.meal-account-name{color:#fff;font-size:.9rem;font-weight:700;line-height:1.3;overflow-wrap:anywhere}
.meal-account-meta{color:#AFC2D7;font-size:.72rem;margin-top:.24rem}
.meal-nav-heading{color:#8FA7C0;font-size:.64rem;font-weight:800;letter-spacing:.16em;margin:.8rem .25rem .25rem;text-transform:uppercase}
.meal-nav-footer{color:#8FA7C0;font-size:.68rem;line-height:1.5;text-align:center;padding:.45rem .2rem}

[data-testid="stMetric"]{background:#fff;border:1px solid #DCE4EF;border-left:5px solid #D4A017;padding:15px;border-radius:12px;box-shadow:0 2px 8px rgba(18,52,90,.06)}

.stButton>button,.stDownloadButton>button,[data-testid="stFormSubmitButton"]>button{background:#12345A;color:#fff;border:1px solid #12345A;border-radius:8px;font-weight:600;min-height:40px}

.stButton>button:hover,.stDownloadButton>button:hover,[data-testid="stFormSubmitButton"]>button:hover{background:#D4A017;color:#12345A;border-color:#D4A017}

.dbk-header{background:linear-gradient(120deg,#12345A,#1D527D);color:#fff;padding:23px;border-radius:14px;border-bottom:5px solid #D4A017;margin-bottom:18px}

.dbk-header h1{color:#fff!important;margin:0;font-size:29px}.dbk-header p{color:#F4E7B2;margin:8px 0 0;font-size:14px}

.dbk-tagline{color:#F4D16A;font-weight:700;letter-spacing:1px;font-size:11px;margin-bottom:7px}

</style>

""", unsafe_allow_html=True)



# =========================================================

# 3. DATABASE

# =========================================================

def get_connection():

          if DATABASE_URL:
                    return DatabaseConnection(
                              psycopg.connect(
                                        DATABASE_URL,
                                        row_factory=dict_row,
                                        prepare_threshold=None,
                                        connect_timeout=10,
                              ),
                              is_postgres=True,
                    )

          conn = sqlite3.connect(DB_FILE, timeout=30)

          conn.row_factory = sqlite3.Row

          return DatabaseConnection(conn, is_postgres=False)


class DatabaseConnection:
          def __init__(self, connection, is_postgres):
                    self.connection = connection
                    self.is_postgres = is_postgres

          def execute(self, query, parameters=()):
                    if self.is_postgres:
                              query = query.replace("BEGIN IMMEDIATE", "BEGIN").replace("?", "%s")
                    return self.connection.execute(query, parameters)

          def __getattr__(self, name):
                    return getattr(self.connection, name)

          def __enter__(self):
                    return self

          def __exit__(self, exception_type, exception, traceback):
                    try:
                              if exception_type is None:
                                        self.connection.commit()
                              else:
                                        self.connection.rollback()
                    finally:
                              self.connection.close()





def hash_password(password):

          iterations = 600_000
          salt = secrets.token_bytes(16)
          password_hash = hashlib.pbkdf2_hmac(
                    "sha256",
                    password.encode("utf-8"),
                    salt,
                    iterations,
          )
          return f"pbkdf2_sha256${iterations}${salt.hex()}${password_hash.hex()}"


def verify_password(password, saved_hash):
          if saved_hash.startswith("pbkdf2_sha256$"):
                    try:
                              algorithm, iterations_text, salt_hex, digest_hex = saved_hash.split("$", 3)
                              iterations = int(iterations_text)
                              if algorithm != "pbkdf2_sha256" or not 100_000 <= iterations <= 2_000_000:
                                        return False, False
                              expected = hashlib.pbkdf2_hmac(
                                        "sha256",
                                        password.encode("utf-8"),
                                        bytes.fromhex(salt_hex),
                                        iterations,
                              ).hex()
                              return secrets.compare_digest(expected, digest_hex), False
                    except (ValueError, TypeError):
                              return False, False

          legacy_hash = hashlib.sha256(password.encode("utf-8")).hexdigest()
          if secrets.compare_digest(legacy_hash, saved_hash):
                    return True, True
          return False, False





def init_db():

          with get_connection() as conn:

                    if DATABASE_URL:
                              conn.execute("""CREATE TABLE IF NOT EXISTS users(
                                        id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                                        name TEXT NOT NULL,
                                        email TEXT UNIQUE NOT NULL,
                                        password TEXT NOT NULL,
                                        role TEXT NOT NULL DEFAULT 'centre_head',
                                        centre TEXT,
                                        approved INTEGER NOT NULL DEFAULT 0,
                                        created_at TEXT)""")
                              conn.execute("""CREATE TABLE IF NOT EXISTS learners(
                                        id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                                        individual_no TEXT, name TEXT NOT NULL, gender TEXT, age INTEGER,
                                        disability TEXT, nationality TEXT, location TEXT, phone TEXT,
                                        centre TEXT, course TEXT, duration TEXT, education TEXT,
                                        certification TEXT, status TEXT DEFAULT 'Ongoing',
                                        dropout_reason TEXT, year INTEGER, created_at TEXT,
                                        completion_month TEXT, dropout_quarter TEXT, record_no TEXT,
                                        institutional_identifier TEXT, kasi TEXT, remarks TEXT,
                                        sponsor TEXT, school_fee TEXT, payment TEXT,
                                        payment_confirmation_no TEXT, reason TEXT)""")
                              conn.execute("""CREATE TABLE IF NOT EXISTS employment(
                                        id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                                        individual_no TEXT, graduate_name TEXT NOT NULL, gender TEXT,
                                        centre TEXT, course TEXT, graduation_year INTEGER,
                                        employment_status TEXT, employer_or_business TEXT,
                                        job_title TEXT, monthly_income DOUBLE PRECISION, phone TEXT,
                                        follow_up_date TEXT, notes TEXT, created_at TEXT)""")
                              conn.execute("""CREATE TABLE IF NOT EXISTS change_audit(
                                        id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                                        actor_id BIGINT, actor_name TEXT NOT NULL, actor_email TEXT,
                                        actor_role TEXT, actor_centre TEXT, action TEXT NOT NULL,
                                        entity_type TEXT NOT NULL, entity_id BIGINT, entity_label TEXT,
                                        before_data TEXT, after_data TEXT, details TEXT,
                                        created_at TEXT NOT NULL)""")
                              conn.execute("CREATE INDEX IF NOT EXISTS idx_learners_centre ON learners(centre)")
                              conn.execute("CREATE INDEX IF NOT EXISTS idx_learners_individual_no ON learners(individual_no)")
                              conn.execute("CREATE INDEX IF NOT EXISTS idx_change_audit_created_at ON change_audit(created_at)")
                              conn.execute("CREATE INDEX IF NOT EXISTS idx_change_audit_entity ON change_audit(entity_type, entity_id)")
                              return

                    conn.execute("""CREATE TABLE IF NOT EXISTS users(

                              id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,

                              email TEXT UNIQUE NOT NULL, password TEXT NOT NULL,

                              role TEXT NOT NULL DEFAULT 'centre_head', centre TEXT,

                              approved INTEGER NOT NULL DEFAULT 0, created_at TEXT)""")

                    conn.execute("""CREATE TABLE IF NOT EXISTS learners(

                              id INTEGER PRIMARY KEY AUTOINCREMENT, individual_no TEXT, name TEXT NOT NULL,

                              gender TEXT, age INTEGER, disability TEXT, nationality TEXT, location TEXT,

                              phone TEXT, centre TEXT, course TEXT, duration TEXT, education TEXT,

                              certification TEXT, status TEXT DEFAULT 'Ongoing', dropout_reason TEXT,

                              year INTEGER, created_at TEXT)""")

                    conn.execute("""CREATE TABLE IF NOT EXISTS employment(

                              id INTEGER PRIMARY KEY AUTOINCREMENT, individual_no TEXT, graduate_name TEXT NOT NULL,

                              gender TEXT, centre TEXT, course TEXT, graduation_year INTEGER,

                              employment_status TEXT, employer_or_business TEXT, job_title TEXT,

                              monthly_income REAL, phone TEXT, follow_up_date TEXT, notes TEXT,

                              created_at TEXT)""")

                    conn.execute("""CREATE TABLE IF NOT EXISTS change_audit(
                              id INTEGER PRIMARY KEY AUTOINCREMENT,
                              actor_id INTEGER,
                              actor_name TEXT NOT NULL,
                              actor_email TEXT,
                              actor_role TEXT,
                              actor_centre TEXT,
                              action TEXT NOT NULL,
                              entity_type TEXT NOT NULL,
                              entity_id INTEGER,
                              entity_label TEXT,
                              before_data TEXT,
                              after_data TEXT,
                              details TEXT,
                              created_at TEXT NOT NULL)""")

                    # Add any missing learner columns when opening an older database.

                    existing = {r["name"] for r in conn.execute("PRAGMA table_info(learners)").fetchall()}

                    learner_columns = {

                              "individual_no":"TEXT", "name":"TEXT", "gender":"TEXT", "age":"INTEGER",

                              "disability":"TEXT", "nationality":"TEXT", "location":"TEXT", "phone":"TEXT",

                              "centre":"TEXT", "course":"TEXT", "duration":"TEXT", "education":"TEXT",

                              "certification":"TEXT", "status":"TEXT DEFAULT 'Ongoing'", "dropout_reason":"TEXT",

                              "year":"INTEGER", "completion_month":"TEXT", "dropout_quarter":"TEXT", "created_at":"TEXT",
                              "record_no":"TEXT", "institutional_identifier":"TEXT", "kasi":"TEXT",
                              "remarks":"TEXT", "sponsor":"TEXT", "school_fee":"TEXT", "payment":"TEXT",
                              "payment_confirmation_no":"TEXT", "reason":"TEXT"

                    }

                    for col, definition in learner_columns.items():

                              if col not in existing:

                                        conn.execute(f"ALTER TABLE learners ADD COLUMN {col} {definition}")

                    user_cols = {r["name"] for r in conn.execute("PRAGMA table_info(users)").fetchall()}

                    if "created_at" not in user_cols:

                              conn.execute("ALTER TABLE users ADD COLUMN created_at TEXT")

                    conn.execute("CREATE INDEX IF NOT EXISTS idx_learners_centre ON learners(centre)")

                    conn.execute("CREATE INDEX IF NOT EXISTS idx_learners_individual_no ON learners(individual_no)")

                    conn.execute("CREATE INDEX IF NOT EXISTS idx_change_audit_created_at ON change_audit(created_at)")
                    conn.execute("CREATE INDEX IF NOT EXISTS idx_change_audit_entity ON change_audit(entity_type, entity_id)")

                    conn.commit()



init_db()



# =========================================================

# 4. HELPERS

# =========================================================

def record_audit(
          conn,
          actor,
          action,
          entity_type,
          entity_id=None,
          entity_label=None,
          before_data=None,
          after_data=None,
          details=None,
):
          conn.execute(
                    """INSERT INTO change_audit(
                              actor_id, actor_name, actor_email, actor_role, actor_centre,
                              action, entity_type, entity_id, entity_label,
                              before_data, after_data, details, created_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                              actor.get("id"),
                              actor.get("name", "Unknown user"),
                              actor.get("email"),
                              actor.get("role"),
                              actor.get("centre"),
                              action,
                              entity_type,
                              entity_id,
                              entity_label,
                              json.dumps(before_data, ensure_ascii=False, default=str) if before_data is not None else None,
                              json.dumps(after_data, ensure_ascii=False, default=str) if after_data is not None else None,
                              json.dumps(details, ensure_ascii=False, default=str) if details is not None else None,
                              datetime.now().isoformat(timespec="seconds"),
                    ),
          )


def audit_record_values(row):
          return {key: row[key] for key in row.keys()}


def get_user_by_email(email):

          with get_connection() as conn:

                    return conn.execute("SELECT * FROM users WHERE email=?", (email.strip().lower(),)).fetchone()





def load_learners():

          with get_connection() as conn:

                    return read_sql_frame(conn, "SELECT * FROM learners ORDER BY id DESC")





def load_employment():

          with get_connection() as conn:

                    return read_sql_frame(conn, "SELECT * FROM employment ORDER BY id DESC")


def read_sql_frame(conn, query, parameters=()):
          result = conn.execute(query, parameters)
          columns = [
                    column.name if hasattr(column, "name") else column[0]
                    for column in result.description
          ]
          return pd.DataFrame([dict(row) for row in result.fetchall()], columns=columns)





def normalise(value):

          if pd.isna(value):

                    return ""

          return str(value).strip().lower()


def normalise_status(value):
          """Map detailed student-facing labels to stable reporting categories."""
          status = normalise(value)
          if status in ("", "ongoing", "in progress"):
                    return "ongoing"
          if status in ("completed", "complete") or status.startswith("completed in "):
                    return "completed"
          if status in ("dropped", "dropout", "dropped out", "drop out") or status.startswith("dropped out q") or status.startswith("dropped out in q"):
                    return "dropped"
          if status in ("transferred", "transfer"):
                    return "transferred"
          return status


def display_student_status(value, completion_month=None, dropout_quarter=None):
          """Return one of the exact status labels used by the student register."""
          base = normalise_status(value)
          month = str(completion_month or "").strip().title()
          quarter = str(dropout_quarter or "").strip().upper()
          raw = str(value or "").strip()
          low = raw.lower()
          if base == "completed":
                    if month in COMPLETION_MONTHS:
                              return f"Completed in {month}"
                    for allowed_month in COMPLETION_MONTHS:
                              if low == f"completed in {allowed_month.lower()}":
                                        return f"Completed in {allowed_month}"
                    return "Completed"
          if base == "dropped":
                    if quarter in DROPOUT_QUARTERS:
                              return f"Dropped out {quarter}"
                    match = re.search(r"q\s*([1-4])", low)
                    if match:
                              return "Dropped out Q" + match.group(1)
                    return "Dropped out"
          if base == "ongoing":
                    return "Ongoing"
          return raw.title() if raw else "Ongoing"





def safe_int(value):

          if value is None or pd.isna(value) or str(value).strip() == "":

                    return None

          try:

                    return int(float(str(value).strip()))

          except (ValueError, TypeError):

                    return None





def individual_exists(individual_no, conn=None):

          if not individual_no:

                    return False

          if conn is not None:

                    return conn.execute("SELECT id FROM learners WHERE TRIM(individual_no)=? LIMIT 1", (str(individual_no).strip(),)).fetchone() is not None

          with get_connection() as connection:

                    return connection.execute("SELECT id FROM learners WHERE TRIM(individual_no)=? LIMIT 1", (str(individual_no).strip(),)).fetchone() is not None





def get_logo_base64():

          try:

                    return base64.b64encode(LOGO_FILE.read_bytes()).decode("utf-8") if LOGO_FILE.exists() else None

          except OSError:

                    return None





def show_branding():

          left, right = st.columns([1, 5])

          with left:

                    if LOGO_FILE.exists():

                              st.image(str(LOGO_FILE), width=125)

                    else:

                              st.markdown('<div style="background:#fff;border:2px solid #D4A017;border-radius:12px;padding:22px 8px;text-align:center;color:#12345A;font-size:25px;font-weight:800">DBK</div>', unsafe_allow_html=True)

          with right:

                    st.markdown(f"""<div class="dbk-header"><div class="dbk-tagline">MONITORING, EVALUATION, ACCOUNTABILITY AND LEARNING</div><h1>Don Bosco Kakuma</h1><p>MEAL Monitoring Dashboard | {datetime.now():%d %B %Y}</p></div>""", unsafe_allow_html=True)





STUDENT_FORMAT_COLUMNS = [
          "No", "Institutional Identifier", "Name", "Gender", "Age", "Disability",
          "Nationality", "Location", "Telephone No", "Individial No",
          "Highest Education Level", "Current Course", "KASI", "Duration",
          "Certification", "Remarks", "Sponsor", "School fee", "Payment",
          "Payment Confirmation No", "Reason"
]


def csv_template():
          return pd.DataFrame(columns=STUDENT_FORMAT_COLUMNS).to_csv(index=False).encode("utf-8")


def excel_template():
          output = BytesIO()
          with pd.ExcelWriter(output, engine="openpyxl") as writer:
                    pd.DataFrame(columns=STUDENT_FORMAT_COLUMNS).to_excel(writer, index=False, sheet_name="Student Records")
                    ws = writer.sheets["Student Records"]
                    ws.freeze_panes = "A2"
                    ws.auto_filter.ref = ws.dimensions
                    for column_cells in ws.columns:
                              letter = column_cells[0].column_letter
                              ws.column_dimensions[letter].width = min(max(len(str(column_cells[0].value or "")) + 3, 14), 30)
          return output.getvalue()


def show_student_template_downloads(key_prefix="student_template"):
          """Offer the same approved student-register template wherever student data is handled."""
          st.caption("Use the same Don Bosco student-register template across registration, uploads, records and reporting.")
          left, right = st.columns(2)
          with left:
                    st.download_button(
                              "Download CSV Student Template",
                              data=csv_template(),
                              file_name="DBK_Student_Register_Template.csv",
                              mime="text/csv",
                              key=f"{key_prefix}_csv",
                              use_container_width=True,
                    )
          with right:
                    st.download_button(
                              "Download Excel Student Template",
                              data=excel_template(),
                              file_name="DBK_Student_Register_Template.xlsx",
                              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              key=f"{key_prefix}_xlsx",
                              use_container_width=True,
                    )


def parse_student_remarks(remarks):
          """Translate the provided format's Remarks into normalized status and timing."""
          value = str(remarks or "").strip()
          low = value.lower()
          status = "Ongoing"
          completion_month = None
          dropout_quarter = None
          if "drop" in low and "q" in low:
                    status = "Dropped"
                    match = re.search(r"q\s*([1-4])", low)
                    if match:
                              dropout_quarter = "Q" + match.group(1)
          else:
                    for month in COMPLETION_MONTHS:
                              if "complet" in low and month.lower() in low:
                                        status = "Completed"
                                        completion_month = month
                                        break
                    if low in ("completed", "complete"):
                              status = "Completed"
          return status, completion_month, dropout_quarter


def student_format_export(frame):
          """Return database records using the user's student-register column names."""
          data = frame.copy()
          def col(name, default=""):
                    if name not in data.columns:
                              return default
                    result = data[name].copy()
                    if isinstance(default, pd.Series):
                              blank = result.isna() | result.astype(str).str.strip().isin(["", "nan", "None"])
                              result = result.where(~blank, default)
                    return result
          if "remarks" in data.columns:
                    remarks = data["remarks"].fillna("").astype(str)
          else:
                    remarks = pd.Series([""] * len(data), index=data.index)
          if "status" in data.columns:
                    for idx in data.index:
                              if not str(remarks.loc[idx]).strip():
                                        status = normalise_status(data.at[idx, "status"])
                                        if status == "completed":
                                                  raw_month = data.at[idx, "completion_month"] if "completion_month" in data.columns else ""
                                                  month = "" if pd.isna(raw_month) else str(raw_month).strip()
                                                  remarks.loc[idx] = f"Completed in {month}" if month else "Completed"
                                        elif status == "dropped":
                                                  raw_quarter = data.at[idx, "dropout_quarter"] if "dropout_quarter" in data.columns else ""
                                                  quarter = "" if pd.isna(raw_quarter) else str(raw_quarter).strip()
                                                  remarks.loc[idx] = f"Dropped out {quarter}" if quarter else "Dropped out"
          out = pd.DataFrame(index=data.index)
          out["No"] = col("record_no")
          out["Institutional Identifier"] = col("institutional_identifier", col("centre"))
          out["Name"] = col("name")
          out["Gender"] = col("gender")
          out["Age"] = col("age")
          out["Disability"] = col("disability")
          out["Nationality"] = col("nationality")
          out["Location"] = col("location")
          out["Telephone No"] = col("phone")
          out["Individial No"] = col("individual_no")
          out["Highest Education Level"] = col("education")
          out["Current Course"] = col("course")
          out["KASI"] = col("kasi")
          out["Duration"] = col("duration")
          out["Certification"] = col("certification")
          out["Remarks"] = remarks
          out["Sponsor"] = col("sponsor")
          out["School fee"] = col("school_fee")
          out["Payment"] = col("payment")
          out["Payment Confirmation No"] = col("payment_confirmation_no")
          out["Reason"] = col("reason", col("dropout_reason"))
          return out



# =========================================================

# 5. EXCEL REPORT GENERATOR

# =========================================================

def category_for_course(course, centre=None):

          c = normalize_trade_text(course)
          centre_name = normalize_trade_text(centre)

          if centre_name == normalize_trade_text("DB Kalobeyei Buss"):

                    return "VOCATIONAL COURSES"

          if any(x in c for x in ["atd", "cams", "cpa", "kasneb", "foundation"]):

                    return "PROFESSIONAL COURSES"

          if "ict" in c:

                    return "ICT NITA COURSE"

          if any(x in c for x in [
                    "masonry", "plumbing", "electrical", "carpentry", "welding",
                    "solar", "lvm", "ppf", "fitting", "turner",
          ]):

                    return "TECHNICAL COURSES"

          if any(x in c for x in ["dress making", "dressmaking", "tailoring"]):

                    if centre_name == normalize_trade_text("DBK"):

                              return "VOCATIONAL COURSES"

                    return "TECHNICAL COURSES"

          if any(x in c for x in [
                    "computer", "secretarial", "fashion", "english", "agriculture", "fish farming",
                    "kitchen gardening", "basketry", "beadwork",
          ]):

                    return "VOCATIONAL COURSES"

          return "PROFESSIONAL COURSES"





def style_and_append(ws, title, headers, rows):

          """Append an automatically generated section below existing template data."""

          start = ws.max_row + 2

          ws.cell(start, 1, title)

          for col in range(1, max(1, len(headers)) + 1):

                    cell = ws.cell(start, col)

                    cell.fill = PatternFill("solid", fgColor="12345A")

                    cell.font = Font(bold=True, color="FFFFFF", size=12)

          header_row = start + 1

          for col, heading in enumerate(headers, 1):

                    cell = ws.cell(header_row, col, heading)

                    cell.fill = PatternFill("solid", fgColor="1D527D")

                    cell.font = Font(bold=True, color="FFFFFF")

                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

          for row_idx, values in enumerate(rows, header_row + 1):

                    for col_idx, value in enumerate(values, 1):

                              if pd.isna(value) if not isinstance(value, (list, dict)) else False:

                                        value = None

                              cell = ws.cell(row_idx, col_idx, value)

                              cell.alignment = Alignment(vertical="top", wrap_text=True)

          for col in range(1, len(headers) + 1):

                    letter = ws.cell(header_row, col).column_letter

                    ws.column_dimensions[letter].width = min(max(len(str(headers[col-1])) + 4, 14), 30)





def as_excel_rows(frame):

          safe = frame.copy().astype(object).where(pd.notna(frame), None)

          return safe.values.tolist()





def build_summary(frame):

          if frame.empty:

                    return pd.DataFrame(columns=["Centre", "Total", "Male", "Female", "Ongoing", "Completed", "Dropped", "Transferred"])

          data = frame.copy()

          data["gender_clean"] = data["gender"].apply(normalise)

          data["status_clean"] = data["status"].apply(normalise_status)

          data["centre"] = data["centre"].fillna("Not specified").replace("", "Not specified")

          output = []

          for centre, group in data.groupby("centre", dropna=False):

                    output.append({

                              "Centre": centre, "Total": len(group),

                              "Male": int((group.gender_clean == "male").sum()),

                              "Female": int((group.gender_clean == "female").sum()),

                              "Ongoing": int((group.status_clean == "ongoing").sum()),

                              "Completed": int((group.status_clean == "completed").sum()),

                              "Dropped": int((group.status_clean == "dropped").sum())

                    })

          return pd.DataFrame(output)





def build_course_summary(frame):

          if frame.empty:

                    return pd.DataFrame(columns=["Course", "Total", "Male", "Female", "Ongoing", "Completed", "Dropped", "Transferred"])

          data = frame.copy()

          data["gender_clean"] = data["gender"].apply(normalise)

          data["status_clean"] = data["status"].apply(normalise_status)

          data["course"] = data["course"].fillna("Not specified").replace("", "Not specified")

          output = []

          for course, group in data.groupby("course", dropna=False):

                    output.append({

                              "Course": course, "Total": len(group),

                              "Male": int((group.gender_clean == "male").sum()),

                              "Female": int((group.gender_clean == "female").sum()),

                              "Ongoing": int((group.status_clean == "ongoing").sum()),

                              "Completed": int((group.status_clean == "completed").sum()),

                              "Dropped": int((group.status_clean == "dropped").sum())

                    })

          return pd.DataFrame(output)


def build_nationality_summary(frame):
          columns = [
                    "Student Status", "Gender", *REPORT_NATIONALITY_ALIASES.keys(), "Total",
          ]
          data = frame.copy()
          if data.empty:
                    data = pd.DataFrame(columns=["nationality", "gender", "status"])
          data["nationality_clean"] = data["nationality"].fillna("").astype(str).apply(normalise)
          code_lookup = {
                    normalise(alias): code
                    for code, aliases in REPORT_NATIONALITY_ALIASES.items()
                    for alias in aliases
          }
          data["nationality_code"] = data["nationality_clean"].map(code_lookup).fillna("Other / Not specified")
          data["gender_clean"] = data["gender"].apply(normalise)
          data["gender_label"] = data["gender_clean"].map(
                    {"male": "Male", "female": "Female"}
          ).fillna("Not specified")
          data["status_clean"] = data["status"].apply(normalise_status)

          statuses = ["ongoing", "dropped", "completed", "transferred"]
          statuses.extend(sorted(set(data["status_clean"]) - set(statuses) - {""}))
          status_labels = {
                    "ongoing": "Ongoing",
                    "dropped": "Dropped",
                    "completed": "Completed",
                    "transferred": "Transferred",
          }
          output = []
          for status in statuses:
                    for gender in ["Male", "Female", "Not specified"]:
                              group = data[
                                        (data["status_clean"] == status)
                                        & (data["gender_label"] == gender)
                              ]
                              row = {
                                        "Student Status": status_labels.get(status, status.title()),
                                        "Gender": gender,
                              }
                              counts = group["nationality_code"].value_counts()
                              row.update({code: int(counts.get(code, 0)) for code in REPORT_NATIONALITY_ALIASES})
                              row["Total"] = len(group)
                              output.append(row)

          for gender in ["Male", "Female", "Not specified"]:
                    group = data[data["gender_label"] == gender]
                    row = {"Student Status": "All Students", "Gender": gender}
                    counts = group["nationality_code"].value_counts()
                    row.update({code: int(counts.get(code, 0)) for code in REPORT_NATIONALITY_ALIASES})
                    row["Total"] = len(group)
                    output.append(row)
          row = {"Student Status": "All Students", "Gender": "All"}
          counts = data["nationality_code"].value_counts()
          row.update({code: int(counts.get(code, 0)) for code in REPORT_NATIONALITY_ALIASES})
          row["Total"] = len(data)
          output.append(row)
          return pd.DataFrame(output, columns=columns)


def normalize_trade_text(value):
          return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", normalise(value))).strip()


def parse_report_trade(trade):
          prefixes = [
                    ("KALOBEYEI BUSINESS CENTRE", "DB Kalobeyei Buss", True),
                    ("KALOBEIYEI BUSINESS CENTRE", "DB Kalobeyei Buss", True),
                    ("KALOBEYEI", "DB Kalobeyei ICT", False),
                    ("KALOBEIYEI", "DB Kalobeyei ICT", False),
                    ("DB 1", "DB1", False),
                    ("DB1", "DB1", False),
                    ("DBI", "DB1", False),
                    ("DB 2", "DB2", False),
                    ("DB2", "DB2", False),
                    ("DB 3", "DB3", False),
                    ("DB3", "DB3", False),
                    ("DB 4", "DB4", False),
                    ("DB4", "DB4", False),
                    ("DBK", "DBK", False),
          ]
          label = str(trade or "").strip()
          for prefix, centre, centre_wide in prefixes:
                    match = re.match(rf"^{re.escape(prefix)}(?:\s+|$)", label, flags=re.IGNORECASE)
                    if match:
                              course_label = normalize_trade_text(label[match.end():])
                              if centre_wide:
                                        return centre, None, True
                              if course_label.startswith("ict nita "):
                                        course_label = "ict " + course_label[len("ict nita "):]
                              if course_label in {
                                        "agriculture and fish farming",
                                        "kitchen gardening and sustainable agriculture",
                              }:
                                        course_keys = {"agriculture"}
                              elif course_label == "secretarial 6 months":
                                        course_keys = {"secretarial studies"}
                              elif course_label == "computer 6 months":
                                        course_keys = {"computer basic"}
                              elif course_label == "fashion and design tailoring iii":
                                        course_keys = {"fashion", "tailoring iii", "dress making iii"}
                              elif course_label.startswith("carpentry and joinery "):
                                        course_keys = {"carpentry " + course_label.rsplit(" ", 1)[-1]}
                              elif course_label.startswith("dressmaking "):
                                        course_keys = {"dress making " + course_label.rsplit(" ", 1)[-1]}
                              elif course_label == "cpa intermaidiate":
                                        course_keys = {"cpa intermediate"}
                              else:
                                        course_keys = {course_label}
                              return centre, course_keys, False

          course_label = normalize_trade_text(label)
          return None, {course_label}, False


def populate_active_trade_sheet(ws, frame, collect_unmatched=False):
          if normalize_trade_text(ws.cell(3, 1).value) != "trade" or normalize_trade_text(ws.cell(3, 2).value) != "gender":
                    return []

          column_codes = {}
          for column in range(3, 17):
                    first_header = normalize_trade_text(ws.cell(3, column).value)
                    second_header = normalize_trade_text(ws.cell(4, column).value)
                    if first_header == "kenya" and second_header == "others":
                              code = "kenya others"
                    elif second_header == "host":
                              code = "host"
                    else:
                              code = first_header
                    code_lookup = {
                              "sud": "Sud", "s sud": "S.Sud", "som": "Som", "ethi": "Ethi",
                              "cong": "Cong", "bur": "Bur", "uga": "Uga", "yemen": "Yemen",
                              "rwan": "Rwan", "afg": "Afg", "mdg": "Mag", "madg": "Mag", "eritrea": "Eritr",
                              "kenya others": "Kenya Others", "host": "Host",
                    }
                    if code in code_lookup:
                              column_codes[code_lookup[code]] = column
          missing_codes = set(REPORT_NATIONALITY_ALIASES) - {"Other / Not specified"} - set(column_codes)
          if missing_codes:
                    raise ValueError(f"Report sheet {ws.title!r} is missing nationality columns: {', '.join(sorted(missing_codes))}")

          code_lookup = {
                    normalise(alias): code
                    for code, aliases in REPORT_NATIONALITY_ALIASES.items()
                    for alias in aliases
          }
          active_records = []
          for _, learner in frame.iterrows():
                    if normalise_status(learner.get("status")) != "ongoing":
                              continue
                    gender = normalise(learner.get("gender"))
                    nationality_code = code_lookup.get(normalise(learner.get("nationality")))
                    if gender not in ("male", "female") or nationality_code not in column_codes:
                              if collect_unmatched:
                                        active_records.append({
                                                  "id": learner.get("id"),
                                                  "name": learner.get("name"),
                                                  "centre": learner.get("centre"),
                                                  "course": learner.get("course"),
                                                  "gender": learner.get("gender"),
                                                  "nationality": learner.get("nationality"),
                                                  "reason": "Gender or nationality is missing or not in the report categories.",
                                        })
                              continue
                    active_records.append({
                              "id": learner.get("id"),
                              "name": learner.get("name"),
                              "centre": learner.get("centre"),
                              "course": learner.get("course"),
                              "gender": gender,
                              "nationality_code": nationality_code,
                    })

          trade_rows = []
          current_trade = None
          for row_number in range(5, ws.max_row + 1):
                    trade_value = ws.cell(row_number, 1).value
                    gender_value = normalise(ws.cell(row_number, 2).value)
                    if trade_value and normalize_trade_text(trade_value) == "grand total":
                              current_trade = None
                              continue
                    if trade_value:
                              current_trade = str(trade_value).strip()
                    if gender_value not in ("m", "f") or not current_trade:
                              continue
                    centre, course_keys, centre_wide = parse_report_trade(current_trade)
                    trade_rows.append({
                              "row": row_number,
                              "gender": "male" if gender_value == "m" else "female",
                              "centre": centre,
                              "course_keys": course_keys,
                              "centre_wide": centre_wide,
                              "trade": current_trade,
                              "counts": {code: 0 for code in column_codes},
                         })

          unmatched = []
          supplemental_counts = {}
          for learner in active_records:
                    if "reason" in learner:
                              unmatched.append(learner)
                              continue
                    if (
                              ws.title in {
                                        "PROFESSIONAL COURSES",
                                        "VOCATIONAL COURSES",
                                        "TECHNICAL COURSES",
                                        "ICT NITA COURSE",
                              }
                              and category_for_course(learner["course"], learner["centre"]) != ws.title
                    ):
                              continue
                    learner_centre = normalize_trade_text(learner["centre"])
                    learner_course = normalize_trade_text(learner["course"])
                    matching_rows = [
                              trade_row for trade_row in trade_rows
                              if trade_row["gender"] == learner["gender"]
                              and (trade_row["centre"] is None or normalize_trade_text(trade_row["centre"]) == learner_centre)
                              and (
                                        trade_row["centre_wide"]
                                        or learner_course in trade_row["course_keys"]
                              )
                    ]
                    if not matching_rows:
                              learner_category = category_for_course(
                                        learner["course"], learner["centre"]
                              )
                              if (
                                        ws.title == "SEPTEMBER STATISTICAL ANALYSIS"
                                        or learner_category == ws.title
                              ):
                                        centre = str(learner["centre"] or "Not specified").strip()
                                        course = str(learner["course"] or "Not specified").strip()
                                        key = (
                                                  normalize_trade_text(centre),
                                                  normalize_trade_text(course),
                                                  learner["gender"],
                                        )
                                        if key not in supplemental_counts:
                                                  supplemental_counts[key] = {
                                                            "trade": f"{centre} {course}".strip(),
                                                            "gender": learner["gender"],
                                                            "counts": {code: 0 for code in column_codes},
                                                  }
                                        supplemental_counts[key]["counts"][
                                                  learner["nationality_code"]
                                        ] += 1
                                        continue
                              if collect_unmatched:
                                        learner["reason"] = "No matching Trade row in the September template."
                                        unmatched.append(learner)
                              continue
                    trade_row = matching_rows[0]
                    trade_row["counts"][learner["nationality_code"]] += 1

          for trade_row in trade_rows:
                    for code, column in column_codes.items():
                              ws.cell(trade_row["row"], column, trade_row["counts"][code])

          gender_totals = {
                    gender: {
                              code: sum(
                                        trade_row["counts"][code]
                                        for trade_row in trade_rows
                                        if trade_row["gender"] == gender
                              )
                              for code in column_codes
                    }
                    for gender in ("male", "female")
          }
          for supplemental in supplemental_counts.values():
                    for code, count in supplemental["counts"].items():
                              gender_totals[supplemental["gender"]][code] += count

          for trade_row in trade_rows:
                    row_number = trade_row["row"]
                    row_total = sum(trade_row["counts"].values())
                    ws.cell(row_number, 17, row_total)
                    if trade_row["gender"] == "male":
                              female_row = next(
                                        (
                                                  row
                                                  for row in trade_rows
                                                  if row["row"] == row_number + 1 and row["gender"] == "female"
                                        ),
                                        None,
                              )
                              ws.cell(
                                        row_number,
                                        18,
                                        row_total + (sum(female_row["counts"].values()) if female_row else 0),
                              )

          grand_total_row = next(
                    (
                              row_number
                              for row_number in range(5, ws.max_row + 1)
                              if normalize_trade_text(ws.cell(row_number, 1).value) == "grand total"
                    ),
                    None,
          )
          if grand_total_row is not None:
                    male_row = grand_total_row
                    female_row = grand_total_row + 1
                    total_row = next(
                              (
                                        row_number
                                        for row_number in range(female_row + 1, ws.max_row + 1)
                                        if normalise(ws.cell(row_number, 2).value) == "t"
                              ),
                              None,
                    )
                    for code, column in column_codes.items():
                              male = gender_totals["male"][code]
                              female = gender_totals["female"][code]
                              ws.cell(male_row, column, male)
                              ws.cell(female_row, column, female)
                              if total_row is not None:
                                        ws.cell(total_row, column, male + female)
                    male_total = sum(gender_totals["male"].values())
                    female_total = sum(gender_totals["female"].values())
                    ws.cell(male_row, 17, male_total)
                    ws.cell(female_row, 17, female_total)
                    if total_row is not None:
                              ws.cell(total_row, 17, male_total + female_total)
                    ws.cell(male_row, 18, male_total + female_total)

                    host_codes = {"Kenya Others", "Host"}
                    for row_number in range(1, ws.max_row + 1):
                              population = normalize_trade_text(ws.cell(row_number, 4).value)
                              if population not in ("refugees", "host", "total"):
                                        continue
                              if population == "refugees":
                                        male = sum(
                                                  value for code, value in gender_totals["male"].items()
                                                  if code not in host_codes
                                        )
                                        female = sum(
                                                  value for code, value in gender_totals["female"].items()
                                                  if code not in host_codes
                                        )
                              elif population == "host":
                                        male = sum(
                                                  value for code, value in gender_totals["male"].items()
                                                  if code in host_codes
                                        )
                                        female = sum(
                                                  value for code, value in gender_totals["female"].items()
                                                  if code in host_codes
                                        )
                              else:
                                        male = male_total
                                        female = female_total
                              ws.cell(row_number, 5, male)
                              ws.cell(row_number, 6, female)
                              ws.cell(row_number, 7, male + female)

          if supplemental_counts:
                    supplemental_rows = [
                              [
                                        supplemental["trade"],
                                        supplemental["gender"].title(),
                                        *[
                                                  supplemental["counts"][code]
                                                  for code in column_codes
                                        ],
                                        sum(supplemental["counts"].values()),
                              ]
                              for supplemental in supplemental_counts.values()
                    ]
                    style_and_append(
                              ws,
                              "ADDITIONAL ACTIVE COURSE DATA NOT LISTED IN TEMPLATE",
                              ["Trade", "Gender", *column_codes.keys(), "Total"],
                              supplemental_rows,
                    )
          return unmatched


def classify_enrollment_course(course, centre):
          course_name = normalize_trade_text(course)
          if normalise(centre) == normalise("DB Kalobeyei Buss"):
                    return 18, "BUSS"
          technical_families = (
                    "electrical", "masonry", "plumbing", "carpentry", "welding",
                    "tailoring", "dress making", "dressmaking", "lvm", "ppf",
                    "solar", "fitting", "turner",
          )
          if any(family in course_name for family in technical_families):
                    level_match = re.search(r"\b(iii|ii|i|3|2|1)\b", course_name)
                    if level_match:
                              level = {"i": 1, "ii": 2, "iii": 3}.get(
                                        level_match.group(1), safe_int(level_match.group(1))
                              )
                              if level in (1, 2, 3):
                                        return 4 + level, f"GRADE {['I', 'II', 'III'][level - 1]}"
                    return None, None
          if "secretarial" in course_name:
                    return 8, "VOCATIONAL"
          if "computer" in course_name:
                    return 9, "ICT"
          if "english intermediate" in course_name:
                    return 10, "VOCATIONAL"
          if "english basic" in course_name:
                    return 11, "VOCATIONAL"
          if any(term in course_name for term in ("agriculture", "fish farming", "agric")):
                    return 12, "VOCATIONAL"
          if any(term in course_name for term in ("fashion", "basketry", "beadwork")):
                    return 13, "VOCATIONAL"
          if course_name in ("kdii", "kdi ii", "kdi 2", "kakuma digital ii"):
                    return 14, "KDII"
          if "kitchen gardening" in course_name:
                    return 15, "VOCATIONAL"
          if any(term in course_name for term in ("kasneb", "atd", "cams", "cpa")):
                    return 16, "PROFESSIONAL"
          if "ict" in course_name:
                    return 17, "ICT"
          return None, None


def populate_enrollment_summary(ws, frame):
          if normalize_trade_text(ws["C3"].value) != "summary data 2026 september enrolment unhcr":
                    return []

          row_counts = {
                    row: {"male": 0, "female": 0, "host_male": 0, "host_female": 0, "refugee_male": 0, "refugee_female": 0}
                    for row in range(5, 19)
          }
          trade_groups = {
                    group: {"host_male": 0, "host_female": 0, "refugee_male": 0, "refugee_female": 0}
                    for group in ("GRADE I", "GRADE II", "GRADE III", "VOCATIONAL", "ICT", "PROFESSIONAL", "KDII", "BUSS")
          }
          disability_counts = {"male": 0, "female": 0}
          unmatched = []
          code_lookup = {
                    normalise(alias): code
                    for code, aliases in REPORT_NATIONALITY_ALIASES.items()
                    for alias in aliases
          }

          for _, learner in frame.iterrows():
                    if normalise_status(learner.get("status")) != "ongoing":
                              continue
                    gender = normalise(learner.get("gender"))
                    if gender not in ("male", "female"):
                              unmatched.append({
                                        "id": learner.get("id"),
                                        "name": learner.get("name"),
                                        "centre": learner.get("centre"),
                                        "course": learner.get("course"),
                                        "reason": "Missing or unsupported gender for enrollment summary.",
                              })
                              continue
                    if normalise(learner.get("disability")) in ("yes", "true", "1"):
                              disability_counts[gender] += 1

                    category_row, trade_group = classify_enrollment_course(
                              learner.get("course"), learner.get("centre")
                    )
                    if category_row is None:
                              unmatched.append({
                                        "id": learner.get("id"),
                                        "name": learner.get("name"),
                                        "centre": learner.get("centre"),
                                        "course": learner.get("course"),
                                        "reason": "Course does not match an enrollment category in the September template.",
                              })
                              continue

                    row_counts[category_row][gender] += 1
                    nationality_code = code_lookup.get(normalise(learner.get("nationality")))
                    population = "host" if nationality_code in ("Host", "Kenya Others") else "refugee"
                    row_counts[category_row][f"{population}_{gender}"] += 1
                    trade_groups[trade_group][f"{population}_{gender}"] += 1

          # Main UNHCR course and gender table.
          for row in range(5, 19):
                    male = row_counts[row]["male"]
                    female = row_counts[row]["female"]
                    ws.cell(row, 5, male)
                    ws.cell(row, 6, female)
                    ws.cell(row, 7, male + female)
          male_total = sum(row_counts[row]["male"] for row in row_counts)
          female_total = sum(row_counts[row]["female"] for row in row_counts)
          ws["E19"] = male_total
          ws["F19"] = female_total
          ws["G19"] = male_total + female_total

          # PWD figures.
          ws["J4"] = disability_counts["male"]
          ws["K4"] = disability_counts["female"]
          ws["L4"] = sum(disability_counts.values())

          # Course-level host/refugee statistics.
          host_refugee_rows = {
                    "GRADE I": (4, 5),
                    "GRADE II": (6, 7),
                    "GRADE III": (8, 9),
                    "VOCATIONAL": (10, 11),
                    "ICT": (12, 13),
                    "PROFESSIONAL": (14, 15),
                    "KDII": (16, 17),
                    "BUSS": (18, 19),
          }
          for group, (host_row, refugee_row) in host_refugee_rows.items():
                    for target_row, population in ((host_row, "host"), (refugee_row, "refugee")):
                              male = trade_groups[group][f"{population}_male"]
                              female = trade_groups[group][f"{population}_female"]
                              ws.cell(target_row, 17, male)
                              ws.cell(target_row, 18, female)
                              ws.cell(target_row, 19, male + female)
          host_male = sum(values["host_male"] for values in trade_groups.values())
          host_female = sum(values["host_female"] for values in trade_groups.values())
          refugee_male = sum(values["refugee_male"] for values in trade_groups.values())
          refugee_female = sum(values["refugee_female"] for values in trade_groups.values())
          ws["K8"], ws["L8"], ws["M8"] = refugee_male, refugee_female, refugee_male + refugee_female
          ws["K9"], ws["L9"], ws["M9"] = host_male, host_female, host_male + host_female
          ws["K10"], ws["L10"], ws["M10"] = host_male + refugee_male, host_female + refugee_female, host_male + host_female + refugee_male + refugee_female
          ws["Q20"], ws["R20"], ws["S20"] = host_male + refugee_male, host_female + refugee_female, host_male + host_female + refugee_male + refugee_female

          # Summary by qualification group.
          group_members = {
                    22: [5, 6, 7],
                    23: [17],
                    24: [8, 9, 10, 11, 12, 13, 14, 15],
                    25: [16],
                    26: [18],
          }
          for target_row, members in group_members.items():
                    male = sum(row_counts[row]["male"] for row in members)
                    female = sum(row_counts[row]["female"] for row in members)
                    ws.cell(target_row, 5, male)
                    ws.cell(target_row, 6, female)
                    ws.cell(target_row, 7, male + female)
          for column in (5, 6):
                    ws.cell(27, column, sum(ws.cell(row, column).value for row in range(22, 27)))
          ws["G27"] = ws["E27"].value + ws["F27"].value

          # The lower table follows the template's vocational, digital-skills and artisan grouping.
          bottom_groups = {
                    31: [5, 6, 7, 8, 10, 11, 12, 13, 14, 15, 16],
                    32: [9, 17],
                    33: [18],
          }
          for target_row, members in bottom_groups.items():
                    male = sum(row_counts[row]["male"] for row in members)
                    female = sum(row_counts[row]["female"] for row in members)
                    ws.cell(target_row, 5, male)
                    ws.cell(target_row, 6, female)
                    ws.cell(target_row, 7, male + female)
          for column in (5, 6):
                    ws.cell(34, column, sum(ws.cell(row, column).value for row in range(31, 34)))
          ws["G34"] = ws["E34"].value + ws["F34"].value

          bottom_host_refugee_rows = {30: ["GRADE I", "GRADE II", "GRADE III", "VOCATIONAL", "PROFESSIONAL", "KDII"], 32: ["ICT"], 34: ["BUSS"]}
          for target_row, groups in bottom_host_refugee_rows.items():
                    for population_row, population in ((target_row, "host"), (target_row + 1, "refugee")):
                              male = sum(trade_groups[group][f"{population}_male"] for group in groups)
                              female = sum(trade_groups[group][f"{population}_female"] for group in groups)
                              ws.cell(population_row, 17, male)
                              ws.cell(population_row, 18, female)
                              ws.cell(population_row, 19, male + female)
          for row in (36,):
                    male = sum(ws.cell(r, 17).value for r in range(30, 36))
                    female = sum(ws.cell(r, 18).value for r in range(30, 36))
                    ws.cell(row, 17, male)
                    ws.cell(row, 18, female)
                    ws.cell(row, 19, male + female)
          return unmatched




def generate_statistical_report(frame, employment_frame):

          if not REPORT_TEMPLATE.exists():

                    raise FileNotFoundError(f"Template not found: {REPORT_TEMPLATE}")

          buffer = BytesIO()

          wb = load_workbook(REPORT_TEMPLATE)
          wb.calculation.calcMode = "auto"
          wb.calculation.fullCalcOnLoad = True
          wb.calculation.forceFullCalc = True

          course_summary = build_course_summary(frame)

          centre_summary = build_summary(frame)
          enrollment_unmatched = []
          if "SEPTEMBER ANALYSIS" in wb.sheetnames:
                    enrollment_unmatched = populate_enrollment_summary(wb["SEPTEMBER ANALYSIS"], frame)

          trade_sheet_names = [
                    "SEPTEMBER STATISTICAL ANALYSIS",
                    "PROFESSIONAL COURSES",
                    "VOCATIONAL COURSES",
                    "TECHNICAL COURSES",
                    "ICT NITA COURSE",
          ]
          unmatched_active_learners = []
          for sheet_name in trade_sheet_names:
                    if sheet_name in wb.sheetnames:
                              unmatched = populate_active_trade_sheet(
                                        wb[sheet_name],
                                        frame,
                                        collect_unmatched=(sheet_name == "SEPTEMBER STATISTICAL ANALYSIS"),
                              )
                              if sheet_name == "SEPTEMBER STATISTICAL ANALYSIS":
                                        unmatched_active_learners = unmatched



          # Update the template's active-training trade rows, then append detailed live summaries.

          summary_sheet_name = next(
                    (name for name in ["SEPTEMBER ANALYSIS", "APRIL ANALYSIS"] if name in wb.sheetnames),
                    wb.sheetnames[0],
          )
          ws = wb[summary_sheet_name]

          gender = frame["gender"].apply(normalise) if not frame.empty else pd.Series(dtype=str)

          status = frame["status"].apply(normalise_status) if not frame.empty else pd.Series(dtype=str)

          summary_rows = [

                    ["Total registered students", len(frame)],

                    ["Male students", int((gender == "male").sum())],

                    ["Female students", int((gender == "female").sum())],

                    ["Ongoing students", int((status == "ongoing").sum())],

                    ["Completed students", int((status == "completed").sum())],

                    ["Dropouts", int((status == "dropped").sum())],

                    ["Generated on", datetime.now().strftime("%d %B %Y %H:%M")],

          ]

          style_and_append(ws, "LIVE MEAL SYSTEM SUMMARY", ["Indicator", "Value"], summary_rows)



          centre_sheet_name = next(
                    (name for name in ["SEPTEMBER STATISTICAL ANALYSIS", "APRIL STATISTICAL ANALYSIS"] if name in wb.sheetnames),
                    None,
          )
          if centre_sheet_name:
                    style_and_append(wb[centre_sheet_name], "LIVE SUMMARY BY CENTRE", list(centre_summary.columns), as_excel_rows(centre_summary))

          nationality_summary = build_nationality_summary(frame)
          style_and_append(
                    wb[summary_sheet_name],
                    "LIVE SUMMARY BY NATIONALITY",
                    list(nationality_summary.columns),
                    as_excel_rows(nationality_summary),
          )
          style_and_append(
                    wb[summary_sheet_name],
                    "NATIONALITY CODE KEY",
                    ["Code", "Nationality"],
                    [
                              ["Sud", "Sudanese"],
                              ["S.Sud", "South Sudanese"],
                              ["Som", "Somali"],
                              ["Ethi", "Ethiopian"],
                              ["Cong", "Congolese"],
                              ["Bur", "Burundian"],
                              ["Kenya Others", "Kenya / Kenyan Other(s)"],
                              ["Host", "Kenya / Kenyan"],
                              ["Uga", "Ugandan"],
                              ["Yemen", "Yemen / Yemeni"],
                              ["Rwan", "Rwandese"],
                              ["Afg", "Afghan"],
                              ["Mag", "Madagasy / Malagasy"],
                              ["Eritr", "Eritrean"],
                              ["Other / Not specified", "Other, blank or unrecognized nationality"],
                    ],
          )

          for sheet_name in ["PROFESSIONAL COURSES", "VOCATIONAL COURSES", "TECHNICAL COURSES", "ICT NITA COURSE"]:

                    if sheet_name in wb.sheetnames:

                              if frame.empty:

                                        rows = []

                              else:

                                        categorized_frame = frame.copy()
                                        categorized_frame["report_category"] = categorized_frame.apply(
                                                  lambda learner: category_for_course(
                                                            learner.get("course"),
                                                            learner.get("centre"),
                                                  ),
                                                  axis=1,
                                        )
                                        category_frame = categorized_frame[
                                                  categorized_frame["report_category"] == sheet_name
                                        ].drop(columns=["report_category"])
                                        subset = build_course_summary(category_frame)

                                        rows = as_excel_rows(subset)

                              style_and_append(wb[sheet_name], "LIVE COURSE STATISTICS", list(course_summary.columns), rows)



          # Completed records using the actual system status.

          completed = frame[frame["status"].apply(normalise_status) == "completed"].copy() if not frame.empty else frame.copy()

          if "COMPLETED" in wb.sheetnames:

                    completed_cols = ["centre", "name", "gender", "age", "disability", "nationality", "individual_no", "course", "certification", "completion_month"]
                    completed = completed.reindex(columns=completed_cols)
                    completed["Remarks"] = completed["completion_month"].apply(lambda m: f"Completed in {m}" if pd.notna(m) and str(m).strip() else "Completed; month not specified")
                    style_and_append(wb["COMPLETED"], "LIVE COMPLETED STUDENTS", ["Remarks", "Centre", "Name", "Gender", "Age", "Disability", "Nationality", "Individual No", "Current Course", "Certification", "Completion Month"], as_excel_rows(completed[["Remarks"] + completed_cols]))



          # Dropouts from database.

          dropped = frame[frame["status"].apply(normalise_status) == "dropped"].copy() if not frame.empty else frame.copy()

          if "DROP OUTS" in wb.sheetnames:

                    dropout_cols = ["centre", "name", "gender", "nationality", "individual_no", "course", "year", "dropout_quarter", "dropout_reason"]
                    dropped = dropped.reindex(columns=dropout_cols)
                    dropped["Remarks"] = dropped["dropout_quarter"].apply(lambda q: f"Dropped out {q}" if pd.notna(q) and str(q).strip() else "Dropped out; quarter not specified")
                    style_and_append(wb["DROP OUTS"], "LIVE DROPOUT RECORDS", ["Remarks", "Centre", "Full Name", "Gender", "Nationality", "Individual No", "Course", "Year/Grade", "Dropout Quarter", "Reason for Dropping"], as_excel_rows(dropped[["Remarks"] + dropout_cols]))



          # Employment records captured using the new Employment Outcomes page.

          if "EMPLOYED" in wb.sheetnames:

                    if employment_frame.empty:

                              employment_rows = [["No employment outcome records have been entered in the MEAL system yet.", "", "", "", "", "", "", "", ""]]

                              employment_headers = ["Notice", "", "", "", "", "", "", "", ""]

                    else:

                              employment_headers = ["Individual No", "Graduate Name", "Gender", "Centre", "Course", "Employment Status", "Employer/Business", "Job Title", "Monthly Income"]

                              employment_rows = as_excel_rows(employment_frame[["individual_no", "graduate_name", "gender", "centre", "course", "employment_status", "employer_or_business", "job_title", "monthly_income"]])

                    style_and_append(wb["EMPLOYED"], "LIVE EMPLOYMENT OUTCOMES", employment_headers, employment_rows)



          # All-years quarterly remarks summary, with status blanks treated as ongoing.
          quarterly_rows = []
          status_norm = frame["status"].apply(normalise_status) if not frame.empty else pd.Series(dtype=str)
          if not frame.empty:
              for month in COMPLETION_MONTHS:
                  count = int(((status_norm == "completed") & (frame["completion_month"].fillna("").astype(str).str.strip().str.lower() == month.lower())).sum())
                  quarterly_rows.append([f"Completed in {month}", count])
              for quarter in DROPOUT_QUARTERS:
                  count = int(((status_norm == "dropped") & (frame["dropout_quarter"].fillna("").astype(str).str.strip().str.upper() == quarter)).sum())
                  quarterly_rows.append([f"Dropped out {quarter}", count])
              quarterly_rows.extend([
                  ["Ongoing (including blank status cells)", int((status_norm == "ongoing").sum())],
                  ["Completed with month not specified", int(((status_norm == "completed") & frame["completion_month"].fillna("").astype(str).str.strip().eq("")).sum())],
                  ["Dropped with quarter not specified", int(((status_norm == "dropped") & frame["dropout_quarter"].fillna("").astype(str).str.strip().eq("")).sum())],
              ])
          else:
              quarterly_rows = [["No learner records available", 0]]
          quarterly_ws = wb["QUARTERLY REMARKS"] if "QUARTERLY REMARKS" in wb.sheetnames else wb.create_sheet("QUARTERLY REMARKS")
          style_and_append(quarterly_ws, "ALL-YEARS COMPLETION, DROPOUT AND ONGOING SUMMARY", ["Indicator / Remark", "Number of Learners"], quarterly_rows)

          if "REPORT INFO" in wb.sheetnames:

                    info = wb["REPORT INFO"]

          else:

                    info = wb.create_sheet("REPORT INFO")

          info.append([])

          for row in [

                    ["Report generated", datetime.now().strftime("%d %B %Y %H:%M")],

                    ["Data source", "Don Bosco Kakuma MEAL database"],

                    ["Student records included", len(frame)],

                    ["Employment records included", len(employment_frame)],

                    ["Ongoing learners not matched to September enrollment categories", len(enrollment_unmatched)],

                    ["Ongoing learners not matched to template trade rows", len(unmatched_active_learners)],

                    ["Important note", "September enrollment and active-training counts use ongoing MEAL learner records. Completed and dropout records are listed in their respective sheets."],

          ]:

                    info.append(row)

          for row in info.iter_rows(min_row=max(1, info.max_row-4), max_row=info.max_row, min_col=1, max_col=2):

                    row[0].font = Font(bold=True)

          if unmatched_active_learners:
                    unmatched_rows = [
                              [
                                        learner.get("id"),
                                        learner.get("name"),
                                        learner.get("centre"),
                                        learner.get("course"),
                                        learner.get("gender"),
                                        learner.get("nationality"),
                                        learner.get("reason"),
                              ]
                              for learner in unmatched_active_learners
                    ]
                    style_and_append(
                              info,
                              "ONGOING LEARNERS NOT MATCHED TO TEMPLATE TRADES",
                              ["Record ID", "Name", "Centre", "Course", "Gender", "Nationality", "Reason"],
                              unmatched_rows,
                    )

          if enrollment_unmatched:
                    unmatched_enrollment_rows = [
                              [
                                        learner.get("id"),
                                        learner.get("name"),
                                        learner.get("centre"),
                                        learner.get("course"),
                                        learner.get("reason"),
                              ]
                              for learner in enrollment_unmatched
                    ]
                    style_and_append(
                              info,
                              "ONGOING LEARNERS NOT MATCHED TO SEPTEMBER ENROLLMENT CATEGORIES",
                              ["Record ID", "Name", "Centre", "Course", "Reason"],
                              unmatched_enrollment_rows,
                    )

          wb.save(buffer)

          buffer.seek(0)

          return buffer.getvalue()



# =========================================================

# 6. LOGIN / SIGN-UP

# =========================================================

if "user" not in st.session_state:

          st.session_state.user = None

if "page" not in st.session_state:

          st.session_state.page = "Dashboard"





def logout():

          st.session_state.user = None

          st.session_state.page = "Dashboard"

          st.rerun()



if st.session_state.user is None:

          with get_connection() as conn:

                    user_count = conn.execute("SELECT COUNT(*) AS total FROM users").fetchone()["total"]

          brand_column, hero_column = st.columns([0.85, 5], vertical_alignment="center")
          with brand_column:
                    if LOGO_FILE.exists():
                              st.image(str(LOGO_FILE), width=118)
                    else:
                              st.markdown(
                                        '<div style="background:#fff;border:2px solid #D4A017;border-radius:16px;padding:24px 10px;text-align:center;color:#12345A;font-size:26px;font-weight:800">DBK</div>',
                                        unsafe_allow_html=True,
                              )
          with hero_column:
                    st.markdown(
                              f"""
                              <div class="welcome-hero">
                                <div>
                                  <div class="welcome-kicker"><span class="welcome-status"></span> Monitoring workspace · {datetime.now():%d %B %Y}</div>
                                  <div class="welcome-title">Don Bosco Kakuma<br>MEAL Insights</div>
                                  <div class="welcome-copy">A central view of learner registration, education progress and employment outcomes — built to support clear, evidence-led decisions.</div>
                                </div>
                                <div class="welcome-chart" aria-hidden="true"><span></span><span></span><span></span><span></span><span></span><span></span></div>
                              </div>
                              """,
                              unsafe_allow_html=True,
                    )

          st.markdown('<div class="welcome-section-title">Access your workspace</div><div class="welcome-section-copy">Sign in to manage learner information and explore programme data, or create an account to request access.</div>', unsafe_allow_html=True)

          login_tab, signup_tab = st.tabs(["Login", "Sign Up"])

          with login_tab:

                    with st.form("login_form"):

                              email = st.text_input("Email address")

                              password = st.text_input("Password", type="password")

                              submitted = st.form_submit_button("Login", use_container_width=True)

                    if submitted:

                              account = get_user_by_email(email)
                              authenticated, upgrade_legacy_hash = (
                                        verify_password(password, account["password"])
                                        if account
                                        else (False, False)
                              )

                              if account and authenticated:

                                        if upgrade_legacy_hash:
                                                  with get_connection() as conn:
                                                            conn.execute(
                                                                      "UPDATE users SET password=? WHERE id=?",
                                                                      (hash_password(password), account["id"]),
                                                            )
                                                            conn.commit()

                                        if not account["approved"]:

                                                  st.warning("Your account is awaiting administrator approval.")

                                        else:

                                                  st.session_state.user = dict(account)

                                                  st.session_state.page = "Dashboard"

                                                  st.rerun()

                              else:

                                        st.error("Incorrect email or password.")

          with signup_tab:

                    if user_count == 0:

                              st.info("The first registered account becomes the administrator.")

                    with st.form("signup_form"):

                              full_name = st.text_input("Full name")

                              signup_email = st.text_input("Email address")

                              signup_centre = st.selectbox("Centre", CENTRES)

                              signup_password = st.text_input("Create password", type="password")

                              confirm_password = st.text_input("Confirm password", type="password")

                              signup_submitted = st.form_submit_button("Create Account", use_container_width=True)

                    if signup_submitted:

                              if not full_name.strip() or not signup_email.strip():

                                        st.error("Enter your name and email address.")

                              elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", signup_email.strip()):

                                        st.error("Enter a valid email address.")

                              elif len(signup_password) < 8:

                                        st.error("Use a password of at least 8 characters.")

                              elif signup_password != confirm_password:

                                        st.error("The passwords do not match.")

                              elif get_user_by_email(signup_email):

                                        st.error("This email is already registered.")

                              else:

                                        role = "admin" if user_count == 0 else "centre_head"

                                        approved = 1 if user_count == 0 else 0

                                        try:

                                                  with get_connection() as conn:

                                                            conn.execute("INSERT INTO users(name,email,password,role,centre,approved,created_at) VALUES(?,?,?,?,?,?,?)", (full_name.strip(), signup_email.strip().lower(), hash_password(signup_password), role, signup_centre, approved, datetime.now().isoformat(timespec="seconds")))

                                                            conn.commit()

                                                  st.success("Administrator account created. Please log in." if role == "admin" else "Account created. Await administrator approval.")

                                        except DB_INTEGRITY_ERRORS:

                                                  st.error("This email is already registered.")

          st.stop()



# =========================================================

# 7. NAVIGATION AND ACCESS

# =========================================================

user = st.session_state.user

is_admin = user["role"] == "admin"

user_centre = user["centre"]

if not is_admin and user_centre not in CENTRES:

          st.error("Your account has an invalid centre assignment.")

          st.stop()



st.sidebar.markdown(
          """
          <div class="meal-brand">
            <div class="meal-brand-kicker">Don Bosco Kakuma</div>
            <div class="meal-brand-title">MEAL INSIGHTS</div>
            <div class="meal-brand-subtitle">Monitoring workspace</div>
          </div>
          """,
          unsafe_allow_html=True,
)
if LOGO_FILE.exists():

          st.sidebar.image(str(LOGO_FILE), width=140)
account_meta = [escape(user["role"].replace("_", " ").title())]
if not is_admin:
          account_meta.append(escape(user_centre or "Centre not assigned"))
st.sidebar.markdown(
          f"""
          <div class="meal-account">
            <div class="meal-account-name">{escape(user["name"])}</div>
            <div class="meal-account-meta">{" &nbsp;·&nbsp; ".join(account_meta)}</div>
          </div>
          """,
          unsafe_allow_html=True,
)

navigation_groups = [
          ("Overview", [("Dashboard", "dashboard")]),
          (
                    "Learner workspace",
                    [
                              ("Register Learner", "person_add"),
                              ("Upload Student Data", "upload_file"),
                              ("Learner Records", "groups"),
                    ],
          ),
          (
                    "Insights & reporting",
                    [
                              ("Employment Outcomes", "work"),
                              ("Reports", "monitoring"),
                              ("Generate Excel Report", "table_view"),
                    ],
          ),
]
if is_admin:
          navigation_groups.append(
                    (
                              "Administration",
                              [
                                        ("User Approvals", "manage_accounts"),
                                        ("Change History", "history"),
                                        ("Delete All Data", "delete_forever"),
                              ],
                    )
          )

pages = [page_name for _, group in navigation_groups for page_name, _ in group]
if st.session_state.page not in pages:

          st.session_state.page = "Dashboard"

for group_name, group in navigation_groups:
          st.sidebar.markdown(
                    f'<div class="meal-nav-heading">{escape(group_name)}</div>',
                    unsafe_allow_html=True,
          )
          for page_name, icon_name in group:
                    if st.sidebar.button(
                              f":material/{icon_name}: {page_name}",
                              key=f"navigation_{page_name.lower().replace(' ', '_')}",
                              type="primary" if st.session_state.page == page_name else "secondary",
                              use_container_width=True,
                    ):
                              st.session_state.page = page_name
                              st.rerun()

page = st.session_state.page

st.sidebar.divider()

if st.sidebar.button(":material/logout: Sign out", key="navigation_logout", use_container_width=True):

          logout()



# =========================================================

# 8. LOAD DATA

# =========================================================

df_all = load_learners()

employment_all = load_employment()

df = df_all.copy() if is_admin else df_all[df_all["centre"] == user_centre].copy()

employment_df = employment_all.copy() if is_admin else employment_all[employment_all["centre"] == user_centre].copy()

show_branding()



# =========================================================

# 9. DASHBOARD

# =========================================================

if page == "Dashboard":

          st.subheader("Dashboard Overview")

          genders = df["gender"].apply(normalise) if not df.empty else pd.Series(dtype=str)

          statuses = df["status"].apply(normalise_status) if not df.empty else pd.Series(dtype=str)

          metrics = [len(df), int((genders == "male").sum()), int((genders == "female").sum()), int((statuses == "ongoing").sum()), int((statuses == "completed").sum()), int((statuses == "dropped").sum())]

          labels = ["Total Registered", "Male Students", "Female Students", "Ongoing Students", "Completed", "Dropouts"]

          for row in range(2):

                    cols = st.columns(3)

                    for i, col in enumerate(cols):

                              idx = row * 3 + i

                              col.metric(labels[idx], f"{metrics[idx]:,}")

          st.subheader("Learners by Status, Gender and Host/Refugee")
          st.caption("Host learners have a Kenyan nationality (including Kenyan Other); all other nationalities are grouped as refugees.")
          status_categories = ["ongoing", "dropped", "completed", "transferred"]
          status_labels = {
                    "ongoing": "Ongoing",
                    "dropped": "Dropped",
                    "completed": "Completed",
                    "transferred": "Transferred",
          }
          breakdown = df.copy()
          if not breakdown.empty:
                    breakdown["status_group"] = breakdown["status"].apply(normalise_status)
                    breakdown["gender_group"] = breakdown["gender"].apply(normalise)
                    kenya_nationality = breakdown["nationality"].fillna("").astype(str).str.contains(
                              r"\bkeny(?:a|an)\b",
                              case=False,
                              regex=True,
                      )
                    breakdown["population_group"] = kenya_nationality.map({True: "Host", False: "Refugee"})
                    extra_statuses = sorted(
                              status for status in breakdown["status_group"].unique()
                              if status not in status_categories
                    )
                    status_categories.extend(extra_statuses)
          else:
                    breakdown["status_group"] = pd.Series(dtype=str)
                    breakdown["gender_group"] = pd.Series(dtype=str)
                    breakdown["population_group"] = pd.Series(dtype=str)

          demographic_rows = []
          for status_group in status_categories:
                    row = {"Student Status": status_labels.get(status_group, status_group.title())}
                    for population, gender in [
                              ("Host", "male"), ("Host", "female"),
                              ("Refugee", "male"), ("Refugee", "female"),
                    ]:
                              row[f"{population} {'Male' if gender == 'male' else 'Female'}"] = int(
                                        (
                                                  (breakdown["status_group"] == status_group)
                                                  & (breakdown["population_group"] == population)
                                                  & (breakdown["gender_group"] == gender)
                                        ).sum()
                              )
                    row["Total"] = int((breakdown["status_group"] == status_group).sum())
                    demographic_rows.append(row)
          demographic_rows.append({
                    "Student Status": "All Students",
                    "Host Male": sum(row["Host Male"] for row in demographic_rows),
                    "Host Female": sum(row["Host Female"] for row in demographic_rows),
                    "Refugee Male": sum(row["Refugee Male"] for row in demographic_rows),
                    "Refugee Female": sum(row["Refugee Female"] for row in demographic_rows),
                    "Total": len(df),
          })
          st.dataframe(pd.DataFrame(demographic_rows), use_container_width=True, hide_index=True)

          st.divider()

          if df.empty:

                    st.info("No student records available. Register or upload students.")

          centre_counts = df["centre"].fillna("Not specified").replace("", "Not specified").value_counts() if not df.empty else pd.Series(dtype="int64")
          centre_order = CENTRES + [centre for centre in centre_counts.index if centre not in CENTRES]
          centre_counts = centre_counts.reindex(centre_order, fill_value=0)

          left, right = st.columns(2)

          with left:

                    st.subheader("Students by Centre")

                    st.bar_chart(centre_counts)

          with right:

                    st.subheader("Students by Course")

                    if df.empty:
                              st.info("Course totals will appear after student records are registered.")
                    else:
                              st.bar_chart(df["course"].fillna("Not specified").value_counts().head(12))

          if not df.empty:

                    st.subheader("Gender Distribution")

                    st.bar_chart(df["gender"].fillna("Not specified").value_counts())

                    st.subheader("Latest Student Registrations")

                    st.dataframe(df.assign(status=df.apply(lambda row: display_student_status(row.get("status"), row.get("completion_month"), row.get("dropout_quarter")), axis=1))[["individual_no", "name", "gender", "centre", "course", "status", "completion_month", "dropout_quarter"]].head(10), use_container_width=True, hide_index=True)

          st.subheader("Employment Outcomes")

          if employment_df.empty:

                    st.info("No employment outcome records have been entered yet.")

          else:

                    st.metric("Graduates with tracer records", len(employment_df))

                    st.dataframe(employment_df.head(10), use_container_width=True, hide_index=True)



# =========================================================

# 10. REGISTER LEARNER

# =========================================================

elif page == "Register Learner":

          st.subheader("Add Student Record")
          st.caption("The fields below follow the same student register template used for bulk uploads, records and reporting.")
          show_student_template_downloads("register_template")

          with st.form("learner_registration_form"):
                    st.markdown("**Student identification and personal details**")
                    c1, c2, c3 = st.columns(3)
                    with c1:
                              record_no = st.text_input("No")
                              if is_admin:
                                        institutional_identifier = st.selectbox("Institutional Identifier / Centre", CENTRES)
                              else:
                                        institutional_identifier = user_centre
                                        st.text_input("Institutional Identifier / Centre", value=user_centre, disabled=True)
                              learner_name = st.text_input("Name *")
                              gender = st.selectbox("Gender", GENDERS)
                              age = st.number_input("Age", min_value=1, max_value=100, value=18)
                    with c2:
                              disability = st.selectbox("Disability", ["No", "Yes", "Not specified"])
                              nationality = st.selectbox("Nationality", NATIONALITIES)
                              location = st.text_input("Location")
                              phone = st.text_input("Telephone No")
                              individual_no = st.text_input("Individial No")
                    with c3:
                              education = st.selectbox("Highest Education Level", ["None", "No formal education", "Primary", "Secondary", "Craft Certificate", "Artisan", "Trade Test", "National Foundation Certificate", "NTFQ", "KASNEB", "University Certificate", "Diploma", "Degree", "Other"])
                              course = st.text_input("Current Course")
                              kasi = st.selectbox("KASI", ["No", "Yes", "Not specified"])
                              duration = st.text_input("Duration", placeholder="e.g. 6 Months")
                              certification = st.selectbox("Certification", ["NITA", "TVET/CDACC", "KNEC", "KASNEB", "Internal", "Institutional", "Other", "Not specified"])

                    st.markdown("**Progress, sponsorship and payment details**")
                    c4, c5, c6 = st.columns(3)
                    with c4:
                              status = st.selectbox("Student Status", STATUSES, help="Choose the exact completion month or dropout quarter. Ongoing is for learners still enrolled.")
                    with c5:
                              sponsor = st.text_input("Sponsor")
                              school_fee = st.text_input("School fee")
                              payment = st.text_input("Payment", placeholder="e.g. Paid 5000 NITA G III")
                    with c6:
                              payment_confirmation_no = st.text_input("Payment Confirmation No")
                              year = st.number_input("Registration Year", min_value=2000, max_value=2100, value=datetime.now().year)
                              reason = st.text_input("Reason")
                    remarks = st.text_input("Remarks", placeholder="e.g. Completed in August / Dropped out Q1")
                    submitted = st.form_submit_button("Save Student Record", use_container_width=True)

          if submitted:
                    if not learner_name.strip():
                              st.error("Student name is required.")
                    elif individual_no.strip() and individual_exists(individual_no):
                              st.error("This Individual Number is already registered.")
                    else:
                              try:
                                        parsed_status, parsed_month, parsed_quarter = parse_student_remarks(remarks)
                                        final_status = status
                                        final_month = parsed_month
                                        final_quarter = parsed_quarter
                                        # The selected exact status is authoritative; Remarks may supply a detail only when Ongoing was selected.
                                        if status == "Ongoing" and remarks.strip() and ("completed" in remarks.lower() or "drop" in remarks.lower()):
                                                  final_status = parsed_status
                                        elif status.startswith("Completed in "):
                                                  final_status = "Completed"
                                                  final_month = status.replace("Completed in ", "", 1)
                                                  final_quarter = None
                                        elif status.startswith("Dropped out Q"):
                                                  final_status = "Dropped"
                                                  final_quarter = status.replace("Dropped out ", "", 1)
                                                  final_month = None
                                        else:
                                                  final_month = None
                                                  final_quarter = None
                                        with get_connection() as conn:
                                                  result = conn.execute("""INSERT INTO learners(
                                                  record_no,institutional_identifier,individual_no,name,gender,age,disability,nationality,location,phone,centre,course,duration,education,certification,status,dropout_reason,year,completion_month,dropout_quarter,created_at,kasi,remarks,sponsor,school_fee,payment,payment_confirmation_no,reason
                                                  ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING id""", (
                                                  record_no.strip() or None, institutional_identifier, individual_no.strip() or None,
                                                  learner_name.strip(), gender, int(age), disability, nationality, location.strip() or None,
                                                  phone.strip() or None, institutional_identifier, course.strip() or None, duration.strip() or None,
                                                  education, certification, final_status, reason.strip() or None, int(year),
                                                  final_month if final_status == "Completed" and final_month else None,
                                                  final_quarter if final_status == "Dropped" and final_quarter else None,
                                                  datetime.now().isoformat(timespec="seconds"), kasi, remarks.strip() or None,
                                                  sponsor.strip() or None, school_fee.strip() or None, payment.strip() or None,
                                                  payment_confirmation_no.strip() or None, reason.strip() or None))
                                                  inserted_id = result.fetchone()["id"]
                                                  saved_record = conn.execute(
                                                            "SELECT * FROM learners WHERE id=?",
                                                            (inserted_id,),
                                                  ).fetchone()
                                                  record_audit(
                                                            conn,
                                                            user,
                                                            "create",
                                                            "learner",
                                                            inserted_id,
                                                            learner_name.strip(),
                                                            after_data=audit_record_values(saved_record),
                                                  )
                                                  conn.commit()
                                        st.success("Student record saved successfully.")
                                        st.rerun()
                              except DB_ERRORS as e:
                                        st.error(f"Could not save student: {e}")


# =========================================================

# 11. BULK UPLOAD

# =========================================================

elif page == "Upload Student Data":

          st.subheader("Upload Student Records")
          st.write("Use the shared student-register template. Download it, fill in the same column headings, then upload CSV or Excel.")
          show_student_template_downloads("upload_template")
          uploaded = st.file_uploader("Choose student register file", type=["csv", "xlsx"])

          last_import = st.session_state.get("last_student_import_result")
          if last_import and (uploaded is None or last_import["signature"] == hashlib.sha256(uploaded.getvalue()).hexdigest()):
                    st.subheader("Latest Import Result")
                    st.write(f"**File rows:** {last_import['source_rows']} · **Imported:** {last_import['imported']} · **Skipped:** {last_import['skipped']}")
                    if last_import["issues"]:
                              st.dataframe(pd.DataFrame({"Import issues": last_import["issues"]}), use_container_width=True, hide_index=True)
                              st.download_button(
                                        "Download Import Issues (CSV)",
                                        data=pd.DataFrame({"Import issues": last_import["issues"]}).to_csv(index=False).encode("utf-8"),
                                        file_name="DBK_Student_Import_Issues.csv",
                                        mime="text/csv",
                                        key="student_import_issues_download",
                              )

          if uploaded is not None:
                    try:
                              upload_signature = hashlib.sha256(uploaded.getvalue()).hexdigest()
                              up = pd.read_csv(uploaded, dtype=str, keep_default_na=False) if uploaded.name.lower().endswith(".csv") else pd.read_excel(uploaded, dtype=str, keep_default_na=False)
                              up.columns = [str(c).strip() for c in up.columns]
                              column_map = {
                                        "No":"record_no", "Institutional Identifier":"institutional_identifier",
                                        "Name":"name", "Full Name":"name", "Gender":"gender", "Age":"age",
                                        "Disability":"disability", "Nationality":"nationality", "Location":"location",
                                        "Telephone No":"phone", "Telephone Number":"phone", "Phone":"phone",
                                        "Individial No":"individual_no", "Individual No":"individual_no", "Individual Number":"individual_no", "Individual ID":"individual_no",
                                        "Highest Education Level":"education", "Education":"education",
                                        "Current Course":"course", "Course":"course", "KASI":"kasi", "Duration":"duration",
                                        "Certification":"certification", "Remarks":"remarks", "Sponsor":"sponsor",
                                        "School fee":"school_fee", "School Fee":"school_fee", "Payment":"payment",
                                        "Payment Confirmation No":"payment_confirmation_no", "Payment Confirmation Number":"payment_confirmation_no",
                                        "Reason":"reason", "Reason for Dropping":"reason", "Dropout Reason":"reason",
                                        "Centre":"institutional_identifier", "Status":"status", "Completion Month":"completion_month",
                                        "Dropout Quarter":"dropout_quarter", "Year":"year", "Registration Year":"year"
                              }
                              up = up.rename(columns=column_map)
                              if "name" not in up.columns:
                                        st.error("The file must contain the 'Name' column from your student register format.")
                                        st.stop()
                              fields = ["record_no", "institutional_identifier", "name", "gender", "age", "disability", "nationality", "location", "phone", "individual_no", "education", "course", "kasi", "duration", "certification", "remarks", "sponsor", "school_fee", "payment", "payment_confirmation_no", "reason", "status", "completion_month", "dropout_quarter", "year"]
                              for field in fields:
                                        if field not in up.columns:
                                                  up[field] = ""
                              up = up[fields].copy()
                              source_rows = len(up)
                              unnamed_rows = up.index[up["name"].astype(str).str.strip() == ""].tolist()
                              up = up[up["name"].astype(str).str.strip() != ""].copy()
                              # Parse progress from Remarks in the supplied register. A blank remark/status means ongoing.
                              for idx, row in up.iterrows():
                                        parsed_status, parsed_month, parsed_quarter = parse_student_remarks(row.get("remarks", ""))
                                        if not str(row.get("status", "")).strip():
                                                  up.at[idx, "status"] = parsed_status
                                        if not str(row.get("completion_month", "")).strip() and parsed_month:
                                                  up.at[idx, "completion_month"] = parsed_month
                                        if not str(row.get("dropout_quarter", "")).strip() and parsed_quarter:
                                                  up.at[idx, "dropout_quarter"] = parsed_quarter
                              if not is_admin:
                                        up["institutional_identifier"] = user_centre
                              st.write(f"**Rows in file:** {source_rows} · **Rows ready for review:** {len(up)} · **Rows with no name:** {len(unnamed_rows)}")
                              if unnamed_rows:
                                        st.warning("Rows without a Name are not imported. They will be listed in the import issues report.")
                              st.dataframe(up.head(30), use_container_width=True, hide_index=True)
                              if st.button("Import Student Data", type="primary", use_container_width=True):
                                        imported, skipped = 0, len(unnamed_rows)
                                        issues = [f"Row {index + 2}: Name is blank." for index in unnamed_rows]
                                        with get_connection() as conn:
                                                  for index, row in up.iterrows():
                                                            rownum = index + 2
                                                            def issue(message):
                                                                      issues.append(f"Row {rownum}: {message}")
                                                            name = str(row["name"]).strip()
                                                            individual_no = str(row["individual_no"]).strip()
                                                            centre_val = str(row["institutional_identifier"]).strip()
                                                            if centre_val not in CENTRES:
                                                                      skipped += 1; issue(f"Institutional Identifier must match a centre: {centre_val!r}."); continue
                                                            if not is_admin and centre_val != user_centre:
                                                                      skipped += 1; issue(f"Centre heads may only import records for {user_centre!r}."); continue
                                                            gender_val = str(row["gender"]).strip().title()
                                                            if gender_val and gender_val not in GENDERS:
                                                                      skipped += 1; issue("Gender must be Male or Female."); continue
                                                            raw_status = str(row["status"]).strip()
                                                            status_val = raw_status
                                                            completion_month_val = str(row["completion_month"]).strip().title()
                                                            dropout_quarter_val = str(row["dropout_quarter"]).strip().upper()
                                                            if raw_status in STATUS_DETAILS:
                                                                      if raw_status.startswith("Completed in "):
                                                                                status_val = "Completed"
                                                                                completion_month_val = raw_status.replace("Completed in ", "", 1)
                                                                                dropout_quarter_val = ""
                                                                      elif raw_status.startswith("Dropped out Q"):
                                                                                status_val = "Dropped"
                                                                                dropout_quarter_val = raw_status.replace("Dropped out ", "", 1).upper()
                                                                                completion_month_val = ""
                                                            else:
                                                                      status_val = {"": "Ongoing", "Complete":"Completed", "Dropout":"Dropped", "Dropped Out":"Dropped", "In Progress":"Ongoing"}.get(raw_status.title(), raw_status.title())
                                                            if status_val not in ("Ongoing", "Completed", "Dropped", "Transferred"):
                                                                      skipped += 1; issue(f"Invalid status {raw_status!r}. Use one of the eight specified status labels, or Ongoing for a current learner."); continue
                                                            if completion_month_val and completion_month_val not in COMPLETION_MONTHS:
                                                                      skipped += 1; issue("Completion Month must be April, June, August or December."); continue
                                                            if dropout_quarter_val and dropout_quarter_val not in DROPOUT_QUARTERS:
                                                                      skipped += 1; issue("Dropout Quarter must be Q1, Q2, Q3 or Q4."); continue
                                                            if status_val != "Completed": completion_month_val = ""
                                                            if status_val != "Dropped": dropout_quarter_val = ""
                                                            if individual_no and individual_exists(individual_no, conn):
                                                                      skipped += 1; issue(f"Individual Number {individual_no!r} already exists."); continue
                                                            record_no = str(row["record_no"]).strip()
                                                            if not individual_no and record_no:
                                                                      existing_register_row = conn.execute(
                                                                                """SELECT id FROM learners
                                                                                WHERE LOWER(TRIM(COALESCE(record_no,'')))=LOWER(?)
                                                                                AND LOWER(TRIM(COALESCE(name,'')))=LOWER(?)
                                                                                AND LOWER(TRIM(COALESCE(centre,'')))=LOWER(?)
                                                                                LIMIT 1""",
                                                                                (record_no, name, centre_val),
                                                                      ).fetchone()
                                                                      if existing_register_row:
                                                                                skipped += 1
                                                                                issue(
                                                                                          f"Register No {record_no!r} for {name!r} at "
                                                                                          f"{centre_val!r} already exists; no Individual Number was provided."
                                                                                )
                                                                                continue
                                                            age_val = safe_int(row["age"])
                                                            year_val = safe_int(row["year"])
                                                            if str(row["age"]).strip() and (age_val is None or not 1 <= age_val <= 100):
                                                                      skipped += 1; issue("Age must be between 1 and 100."); continue
                                                            if str(row["year"]).strip() and (year_val is None or not 2000 <= year_val <= 2100):
                                                                      skipped += 1; issue("Invalid registration year."); continue
                                                            reason_val = str(row["reason"]).strip()
                                                            result = conn.execute("""INSERT INTO learners(
                                                            individual_no,name,gender,age,disability,nationality,location,phone,centre,course,duration,education,certification,status,dropout_reason,year,completion_month,dropout_quarter,created_at,record_no,institutional_identifier,kasi,remarks,sponsor,school_fee,payment,payment_confirmation_no,reason
                                                            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING id""", (
                                                            individual_no or None, name, gender_val or None, age_val,
                                                            str(row["disability"]).strip() or None, str(row["nationality"]).strip() or None,
                                                            str(row["location"]).strip() or None, str(row["phone"]).strip() or None,
                                                            centre_val, str(row["course"]).strip() or None, str(row["duration"]).strip() or None,
                                                            str(row["education"]).strip() or None, str(row["certification"]).strip() or None,
                                                            status_val, reason_val or None, year_val,
                                                            completion_month_val or None, dropout_quarter_val or None,
                                                            datetime.now().isoformat(timespec="seconds"), record_no or None,
                                                            centre_val, str(row["kasi"]).strip() or None, str(row["remarks"]).strip() or None,
                                                            str(row["sponsor"]).strip() or None, str(row["school_fee"]).strip() or None,
                                                            str(row["payment"]).strip() or None, str(row["payment_confirmation_no"]).strip() or None,
                                                            reason_val or None))
                                                            inserted_id = result.fetchone()["id"]
                                                            saved_record = conn.execute(
                                                                      "SELECT * FROM learners WHERE id=?",
                                                                      (inserted_id,),
                                                            ).fetchone()
                                                            record_audit(
                                                                      conn,
                                                                      user,
                                                                      "import",
                                                                      "learner",
                                                                      inserted_id,
                                                                      name,
                                                                      after_data=audit_record_values(saved_record),
                                                                      details={"file": uploaded.name, "source_row": rownum},
                                                            )
                                                            imported += 1
                                                  record_audit(
                                                            conn,
                                                            user,
                                                            "import_batch",
                                                            "learner_import",
                                                            entity_label=uploaded.name,
                                                            details={
                                                                      "file": uploaded.name,
                                                                      "source_rows": source_rows,
                                                                      "imported": imported,
                                                                      "skipped": skipped,
                                                            },
                                                  )
                                                  conn.commit()
                                        st.session_state["last_student_import_result"] = {
                                                  "signature": upload_signature,
                                                  "source_rows": source_rows,
                                                  "imported": imported,
                                                  "skipped": skipped,
                                                  "issues": issues,
                                        }
                                        st.rerun()
                    except Exception as e:
                              st.error(f"Could not read the uploaded file: {e}")


# =========================================================

# 12. LEARNER RECORDS

# =========================================================

elif page == "Learner Records":

          st.subheader("Student Records")
          show_student_template_downloads("records_template")

          if df.empty:

                    st.info("There are no student records to display.")

          else:

                    a, b = st.columns(2)

                    with a:

                              centre_options = ["All"] + sorted(df["centre"].dropna().astype(str).unique().tolist())

                              selected_centre = st.selectbox("Filter by Centre", centre_options)

                              course_options = ["All"] + sorted(df["course"].dropna().astype(str).unique().tolist())

                              selected_course = st.selectbox("Filter by Course", course_options)

                    with b:

                              selected_gender = st.selectbox("Filter by Gender", ["All"] + GENDERS)

                              selected_status = st.selectbox("Filter by Status", ["All"] + STATUSES + ["Completed", "Dropped out"])

                    search = st.text_input("Search by name or Individual Number")

                    records = df.copy()

                    if selected_centre != "All": records = records[records["centre"] == selected_centre]

                    if selected_course != "All": records = records[records["course"] == selected_course]

                    if selected_gender != "All": records = records[records["gender"] == selected_gender]

                    if selected_status != "All":
                              if selected_status == "Completed":
                                        records = records[records["status"].apply(normalise_status) == "completed"]
                              elif selected_status == "Dropped out":
                                        records = records[records["status"].apply(normalise_status) == "dropped"]
                              else:
                                        records = records[records.apply(lambda row: display_student_status(row.get("status"), row.get("completion_month"), row.get("dropout_quarter")), axis=1) == selected_status]

                    if search.strip():

                              term = search.strip().lower()

                              match_name = records["name"].fillna("").astype(str).str.lower().str.contains(term, regex=False)

                              match_no = records["individual_no"].fillna("").astype(str).str.lower().str.contains(term, regex=False)

                              records = records[match_name | match_no]

                    st.write(f"**Records found: {len(records)}**")

                    st.dataframe(records, use_container_width=True, hide_index=True)

                    export_records = student_format_export(records)
                    st.download_button("Export Filtered Records in Student Register Format (CSV)", export_records.to_csv(index=False).encode("utf-8"), "DBK_Filtered_Student_Records.csv", "text/csv")

                    st.divider()
                    st.subheader("Edit Learner Record")
                    st.caption("Centre heads can edit learner records assigned to their centre and mark ongoing learners as dropped. Dropout reason and quarter are optional.")
                    edit_options = {
                              int(row["id"]): (
                                        f"{row['name']} — {row['centre'] or 'Centre not specified'} — "
                                        f"{row['individual_no'] or 'No Individual Number'}"
                              )
                              for _, row in records.iterrows()
                    }
                    if not edit_options:
                              edit_options = {None: "No learners match the current filters"}
                    selected_edit_id = st.selectbox(
                              "Choose a learner to edit",
                              options=list(edit_options),
                              format_func=lambda record_id: edit_options[record_id],
                              key="edit_learner_record_id",
                    )
                    selected_row = (
                              records[records["id"] == selected_edit_id].iloc[0].to_dict()
                              if selected_edit_id is not None
                              else {"gender": "Male"}
                    )
                    current_status = normalise_status(selected_row.get("status"))
                    status_options = (
                              ["Ongoing", "Dropped"]
                              if current_status != "completed"
                              else ["Completed"]
                    )
                    if is_admin:
                              status_options = ["Ongoing", "Dropped", "Completed", "Transferred"]
                    status_display = {
                              "ongoing": "Ongoing",
                              "dropped": "Dropped",
                              "completed": "Completed",
                              "transferred": "Transferred",
                    }.get(current_status, "Ongoing")
                    if status_display not in status_options:
                              status_options.append(status_display)
                    key_suffix = str(selected_edit_id)
                    with st.form(f"edit_learner_form_{key_suffix}"):
                              left, middle, right = st.columns(3)
                              with left:
                                        edited_record_no = st.text_input("No", value=str(selected_row.get("record_no") or ""))
                                        edited_name = st.text_input("Name *", value=str(selected_row.get("name") or ""))
                                        edited_gender = st.selectbox(
                                                  "Gender",
                                                  GENDERS,
                                                  index=GENDERS.index(str(selected_row.get("gender") or "").title())
                                                  if str(selected_row.get("gender") or "").title() in GENDERS
                                                  else 0,
                                        )
                                        edited_age = st.text_input("Age", value=str(selected_row.get("age") or ""))
                                        edited_disability = st.text_input("Disability", value=str(selected_row.get("disability") or ""))
                                        edited_nationality = st.text_input("Nationality", value=str(selected_row.get("nationality") or ""))
                                        edited_location = st.text_input("Location", value=str(selected_row.get("location") or ""))
                                        edited_phone = st.text_input("Telephone No", value=str(selected_row.get("phone") or ""))
                              with middle:
                                        centre_options = list(dict.fromkeys(CENTRES + [str(selected_row.get("centre") or "")]))
                                        if is_admin:
                                                  edited_centre = st.selectbox(
                                                            "Centre",
                                                            centre_options,
                                                            index=centre_options.index(str(selected_row.get("centre") or "")),
                                                  )
                                        else:
                                                  edited_centre = user_centre
                                                  st.text_input("Centre", value=user_centre, disabled=True)
                                        if is_admin:
                                                  edited_institutional_identifier = st.text_input(
                                                            "Institutional Identifier",
                                                            value=str(selected_row.get("institutional_identifier") or ""),
                                                  )
                                        else:
                                                  edited_institutional_identifier = user_centre
                                                  st.text_input(
                                                            "Institutional Identifier",
                                                            value=user_centre,
                                                            disabled=True,
                                                  )
                                        edited_individual_no = st.text_input(
                                                  "Individual No",
                                                  value=str(selected_row.get("individual_no") or ""),
                                        )
                                        edited_course = st.text_input("Current Course", value=str(selected_row.get("course") or ""))
                                        edited_duration = st.text_input("Duration", value=str(selected_row.get("duration") or ""))
                                        edited_education = st.text_input("Highest Education Level", value=str(selected_row.get("education") or ""))
                                        edited_certification = st.text_input("Certification", value=str(selected_row.get("certification") or ""))
                                        edited_year = st.text_input("Registration Year", value=str(selected_row.get("year") or ""))
                              with right:
                                        edited_status = st.selectbox(
                                                  "Student Status",
                                                  status_options,
                                                  index=status_options.index(status_display),
                                                  disabled=not is_admin and current_status == "completed",
                                        )
                                        edited_dropout_quarter = st.selectbox(
                                                  "Dropout Quarter (optional)",
                                                  ["", *DROPOUT_QUARTERS],
                                                  index=(["", *DROPOUT_QUARTERS].index(str(selected_row.get("dropout_quarter") or "").upper())
                                                         if str(selected_row.get("dropout_quarter") or "").upper() in ["", *DROPOUT_QUARTERS] else 0),
                                        )
                                        edited_completion_month = st.selectbox(
                                                  "Completion Month",
                                                  ["", *COMPLETION_MONTHS],
                                                  index=(["", *COMPLETION_MONTHS].index(str(selected_row.get("completion_month") or "").title())
                                                         if str(selected_row.get("completion_month") or "").title() in ["", *COMPLETION_MONTHS] else 0),
                                        )
                                        edited_dropout_reason = st.text_input(
                                                  "Dropout Reason (optional)",
                                                  value=str(selected_row.get("dropout_reason") or ""),
                                        )
                                        edited_reason = st.text_input("Reason", value=str(selected_row.get("reason") or ""))
                                        edited_kasi = st.text_input("KASI", value=str(selected_row.get("kasi") or ""))
                                        edited_remarks = st.text_input("Remarks", value=str(selected_row.get("remarks") or ""))
                                        edited_sponsor = st.text_input("Sponsor", value=str(selected_row.get("sponsor") or ""))
                                        edited_school_fee = st.text_input("School Fee", value=str(selected_row.get("school_fee") or ""))
                                        edited_payment = st.text_input("Payment", value=str(selected_row.get("payment") or ""))
                                        edited_payment_confirmation = st.text_input(
                                                  "Payment Confirmation No",
                                                  value=str(selected_row.get("payment_confirmation_no") or ""),
                                        )
                              save_learner_changes = st.form_submit_button(
                                        "Save Learner Changes",
                                        type="primary",
                                        use_container_width=True,
                                        disabled=selected_edit_id is None,
                              )
                    if save_learner_changes:
                              age_value = safe_int(edited_age)
                              year_value = safe_int(edited_year)
                              if not edited_name.strip():
                                        st.error("Student name is required.")
                              elif edited_age.strip() and (age_value is None or not 1 <= age_value <= 100):
                                        st.error("Age must be a number from 1 to 100.")
                              elif edited_year.strip() and (year_value is None or not 2000 <= year_value <= 2100):
                                        st.error("Registration year must be between 2000 and 2100.")
                              else:
                                        edited_values = {
                                                  "record_no": edited_record_no.strip() or None,
                                                  "institutional_identifier": edited_institutional_identifier.strip() or None,
                                                  "individual_no": edited_individual_no.strip() or None,
                                                  "name": edited_name.strip(),
                                                  "gender": edited_gender,
                                                  "age": age_value,
                                                  "disability": edited_disability.strip() or None,
                                                  "nationality": edited_nationality.strip() or None,
                                                  "location": edited_location.strip() or None,
                                                  "phone": edited_phone.strip() or None,
                                                  "centre": edited_centre,
                                                  "course": edited_course.strip() or None,
                                                  "duration": edited_duration.strip() or None,
                                                  "education": edited_education.strip() or None,
                                                  "certification": edited_certification.strip() or None,
                                                  "status": edited_status,
                                                  "dropout_reason": edited_dropout_reason.strip() or None,
                                                  "year": year_value,
                                                  "completion_month": edited_completion_month if edited_status == "Completed" else None,
                                                  "dropout_quarter": edited_dropout_quarter if edited_status == "Dropped" else None,
                                                  "kasi": edited_kasi.strip() or None,
                                                  "remarks": edited_remarks.strip() or None,
                                                  "sponsor": edited_sponsor.strip() or None,
                                                  "school_fee": edited_school_fee.strip() or None,
                                                  "payment": edited_payment.strip() or None,
                                                  "payment_confirmation_no": edited_payment_confirmation.strip() or None,
                                                  "reason": edited_reason.strip() or None,
                                        }
                                        editable_columns = tuple(edited_values)
                                        assignments = ", ".join(f"{column}=?" for column in editable_columns)
                                        try:
                                                  with get_connection() as conn:
                                                            conn.execute("BEGIN IMMEDIATE")
                                                            row = conn.execute(
                                                                      "SELECT * FROM learners WHERE id=?",
                                                                      (int(selected_edit_id),),
                                                            ).fetchone()
                                                            if row is None or (
                                                                      not is_admin and row["centre"] != user_centre
                                                            ):
                                                                      raise PermissionError(
                                                                                "This learner is no longer available in your centre."
                                                                      )
                                                            before_values = audit_record_values(row)
                                                            conn.execute(
                                                                      f"UPDATE learners SET {assignments} WHERE id=?",
                                                                      (*[edited_values[column] for column in editable_columns], int(selected_edit_id)),
                                                            )
                                                            after_row = conn.execute(
                                                                      "SELECT * FROM learners WHERE id=?",
                                                                      (int(selected_edit_id),),
                                                            ).fetchone()
                                                            after_values = audit_record_values(after_row)
                                                            if before_values == after_values:
                                                                      conn.rollback()
                                                                      st.info("No learner data changed.")
                                                            else:
                                                                      record_audit(
                                                                                conn,
                                                                                user,
                                                                                "update",
                                                                                "learner",
                                                                                int(selected_edit_id),
                                                                                edited_name.strip(),
                                                                                before_data=before_values,
                                                                                after_data=after_values,
                                                                      )
                                                                      conn.commit()
                                                                      st.success(
                                                                                "Learner details updated and saved in the change history."
                                                                      )
                                                                      st.rerun()
                                        except LEARNER_UPDATE_ERRORS as error:
                                                  st.error(f"Could not update learner record: {error}")

                    if is_admin:
                              st.divider()
                              st.subheader("Delete Student Records")
                              st.warning("Deletion is permanent. This removes only the selected student records; employment outcome records are not affected.")
                              student_options = {
                                        int(row["id"]): (
                                                  f"ID {int(row['id'])} — {row['name']} — "
                                                  f"{row['centre'] or 'Centre not specified'} — "
                                                  f"{row['individual_no'] or 'No Individual Number'}"
                                        )
                                        for _, row in records.iterrows()
                              }
                              selected_student_ids = st.multiselect(
                                        "Student records to delete",
                                        options=list(student_options),
                                        format_func=lambda record_id: student_options[record_id],
                                        key="delete_student_record_ids",
                              )
                              if selected_student_ids:
                                        st.write(f"**Selected records:** {len(selected_student_ids)}")
                                        st.dataframe(
                                                  records[records["id"].isin(selected_student_ids)],
                                                  use_container_width=True,
                                                  hide_index=True,
                                        )
                                        with st.form("delete_student_record_form"):
                                                  confirm_student_delete = st.checkbox(
                                                            f"I confirm permanent deletion of these {len(selected_student_ids)} student record(s)."
                                                  )
                                                  delete_student = st.form_submit_button(
                                                            f"Delete {len(selected_student_ids)} Selected Student Record(s)",
                                                            type="primary",
                                                            use_container_width=True,
                                                  )
                                        if delete_student:
                                                  if not confirm_student_delete:
                                                            st.error("Confirm the deletion before continuing.")
                                                  else:
                                                            placeholders = ",".join("?" for _ in selected_student_ids)
                                                            with get_connection() as conn:
                                                                      conn.execute("BEGIN IMMEDIATE")
                                                                      selected_records = conn.execute(
                                                                                f"SELECT * FROM learners WHERE id IN ({placeholders})",
                                                                                tuple(selected_student_ids),
                                                                      ).fetchall()
                                                                      for selected_record in selected_records:
                                                                                record_audit(
                                                                                          conn,
                                                                                          user,
                                                                                          "delete",
                                                                                          "learner",
                                                                                          selected_record["id"],
                                                                                          selected_record["name"],
                                                                                          before_data=audit_record_values(selected_record),
                                                                                )
                                                                      result = conn.execute(
                                                                                f"DELETE FROM learners WHERE id IN ({placeholders})",
                                                                                tuple(selected_student_ids),
                                                                      )
                                                                      if result.rowcount == len(selected_student_ids):
                                                                                conn.commit()
                                                                      else:
                                                                                conn.rollback()
                                                            if result.rowcount != len(selected_student_ids):
                                                                      st.error(
                                                                                f"Expected to delete {len(selected_student_ids)} records, "
                                                                                f"but deleted {result.rowcount}. Refresh the page and review the remaining records."
                                                                      )
                                                            else:
                                                                      st.success(f"Deleted {result.rowcount} student record(s).")
                                                                      st.rerun()



# =========================================================

# 13. EMPLOYMENT OUTCOMES / GRADUATE TRACER

# =========================================================

elif page == "Employment Outcomes":

          st.subheader("Graduate Tracer Study and Employment Outcomes")

          st.write("Record employment details so the EMPLOYED worksheet can be generated from the MEAL system.")
          st.caption("Student identity and enrolment details use the shared student-register template; employment-specific fields remain in the tracer form below.")
          show_student_template_downloads("employment_template")

          with st.form("employment_form"):

                    left, right = st.columns(2)

                    with left:

                              individual_no = st.text_input("Individual Number (if available)")

                              graduate_name = st.text_input("Graduate Full Name")

                              gender = st.selectbox("Gender", GENDERS)

                              centre = st.selectbox("Centre", CENTRES, disabled=not is_admin, index=CENTRES.index(user_centre) if not is_admin else 0)

                              course = st.text_input("Course Completed")

                              graduation_year = st.number_input("Graduation Year", min_value=2000, max_value=2100, value=datetime.now().year)

                    with right:

                              employment_status = st.selectbox("Employment Status", ["Employed", "Self-employed", "Business owner", "Unemployed", "Further education", "Unknown"])

                              employer = st.text_input("Employer / Business Name")

                              job_title = st.text_input("Job Title / Type of Work")

                              monthly_income = st.number_input("Monthly Income (KES; optional)", min_value=0.0, value=0.0, step=500.0)

                              phone = st.text_input("Telephone Number")

                              follow_up_date = st.date_input("Follow-up Date", value=datetime.now().date())

                    notes = st.text_area("Notes / Employment details")

                    save_employment = st.form_submit_button("Save Employment Record", use_container_width=True)

          if save_employment:

                    if not graduate_name.strip():

                              st.error("Graduate name is required.")

                    else:

                              with get_connection() as conn:

                                        result = conn.execute("""INSERT INTO employment(individual_no,graduate_name,gender,centre,course,graduation_year,employment_status,employer_or_business,job_title,monthly_income,phone,follow_up_date,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?) RETURNING id""", (individual_no.strip() or None, graduate_name.strip(), gender, centre, course.strip(), int(graduation_year), employment_status, employer.strip(), job_title.strip(), float(monthly_income) if monthly_income else None, phone.strip(), follow_up_date.isoformat(), notes.strip(), datetime.now().isoformat(timespec="seconds")))
                                        inserted_id = result.fetchone()["id"]
                                        saved_employment = conn.execute(
                                                  "SELECT * FROM employment WHERE id=?",
                                                  (inserted_id,),
                                        ).fetchone()
                                        record_audit(
                                                  conn,
                                                  user,
                                                  "create",
                                                  "employment",
                                                  inserted_id,
                                                  graduate_name.strip(),
                                                  after_data=audit_record_values(saved_employment),
                                        )
                                        conn.commit()

                              st.success("Employment outcome saved.")

                              st.rerun()

          st.divider()

          st.subheader("Recorded Employment Outcomes")

          if employment_df.empty:

                    st.info("No employment records found for this centre or account.")

          else:

                    st.dataframe(employment_df, use_container_width=True, hide_index=True)

                    st.download_button("Download Employment Records CSV", employment_df.to_csv(index=False).encode("utf-8"), "DBK_Employment_Outcomes.csv", "text/csv")
                    if is_admin:
                              st.divider()
                              st.subheader("Delete Employment Records")
                              st.warning("Deletion is permanent. Select one or more records and confirm before deleting.")
                              employment_options = {
                                        int(row["id"]): (
                                                  f"ID {int(row['id'])} — {row['graduate_name']} — "
                                                  f"{row['centre'] or 'Centre not specified'} — "
                                                  f"{row['course'] or 'Course not specified'} — "
                                                  f"{row['graduation_year'] or 'Year not specified'}"
                                        )
                                        for _, row in employment_df.iterrows()
                              }
                              selected_employment_ids = st.multiselect(
                                        "Employment records to delete",
                                        options=list(employment_options),
                                        format_func=lambda record_id: employment_options[record_id],
                                        key="delete_employment_record_ids",
                              )
                              if selected_employment_ids:
                                        st.write(f"**Selected records:** {len(selected_employment_ids)}")
                                        st.dataframe(
                                                  employment_df[employment_df["id"].isin(selected_employment_ids)],
                                                  use_container_width=True,
                                                  hide_index=True,
                                        )
                                        with st.form("delete_employment_record_form"):
                                                  confirm_employment_delete = st.checkbox(
                                                            f"I confirm permanent deletion of these {len(selected_employment_ids)} employment record(s)."
                                                  )
                                                  delete_employment = st.form_submit_button(
                                                            f"Delete {len(selected_employment_ids)} Selected Employment Record(s)",
                                                            type="primary",
                                                            use_container_width=True,
                                                  )
                                        if delete_employment:
                                                  if not confirm_employment_delete:
                                                            st.error("Confirm the deletion before continuing.")
                                                  else:
                                                            placeholders = ",".join("?" for _ in selected_employment_ids)
                                                            with get_connection() as conn:
                                                                      conn.execute("BEGIN IMMEDIATE")
                                                                      selected_records = conn.execute(
                                                                                f"SELECT * FROM employment WHERE id IN ({placeholders})",
                                                                                tuple(selected_employment_ids),
                                                                      ).fetchall()
                                                                      for selected_record in selected_records:
                                                                                record_audit(
                                                                                          conn,
                                                                                          user,
                                                                                          "delete",
                                                                                          "employment",
                                                                                          selected_record["id"],
                                                                                          selected_record["graduate_name"],
                                                                                          before_data=audit_record_values(selected_record),
                                                                                )
                                                                      result = conn.execute(
                                                                                f"DELETE FROM employment WHERE id IN ({placeholders})",
                                                                                tuple(selected_employment_ids),
                                                                      )
                                                                      if result.rowcount == len(selected_employment_ids):
                                                                                conn.commit()
                                                                      else:
                                                                                conn.rollback()
                                                            if result.rowcount != len(selected_employment_ids):
                                                                      st.error(
                                                                                f"Expected to delete {len(selected_employment_ids)} records, "
                                                                                f"but deleted {result.rowcount}. Refresh the page and review the remaining records."
                                                                      )
                                                            else:
                                                                      st.success(f"Deleted {result.rowcount} employment record(s).")
                                                                      st.rerun()



# =========================================================

# 14. REPORTS

# =========================================================

elif page == "Reports":

          st.subheader("Monitoring Reports")
          show_student_template_downloads("reports_template")

          if df.empty:

                    st.info("Register or upload student records to generate reports.")

          else:

                    st.write("### Summary by Centre")

                    centre_report = build_summary(df)

                    st.dataframe(centre_report, use_container_width=True, hide_index=True)

                    st.write("### Summary by Course")

                    course_report = build_course_summary(df)

                    st.dataframe(course_report, use_container_width=True, hide_index=True)

                    st.write("### Dropout Records")

                    drops = df[df["status"].apply(normalise_status) == "dropped"].copy()

                    if drops.empty:

                              st.info("No dropout records are currently recorded.")

                    else:

                              st.dataframe(student_format_export(drops), use_container_width=True, hide_index=True)

                              st.download_button("Export Dropout Report (Student Register Format)", student_format_export(drops).to_csv(index=False).encode("utf-8"), "DBK_Dropout_Report.csv", "text/csv")

                    st.download_button("Download Full Student Dataset (Student Register Format)", student_format_export(df).to_csv(index=False).encode("utf-8"), "DBK_Student_Dataset.csv", "text/csv")



# =========================================================

# 15. GENERATE EXCEL REPORT

# =========================================================

elif page == "Generate Excel Report":

          st.subheader("Automatic Excel Statistical Report")

          st.write("Generate a downloadable Excel workbook using live records from the MEAL database and the existing statistical report template.")

          st.caption(f"Template expected: {REPORT_TEMPLATE.name}")

          if not REPORT_TEMPLATE.exists():

                    st.error(f"Excel template not found: {REPORT_TEMPLATE}")

          else:

                    genders = df["gender"].apply(normalise) if not df.empty else pd.Series(dtype=str)

                    statuses = df["status"].apply(normalise_status) if not df.empty else pd.Series(dtype=str)

                    cols = st.columns(5)

                    cols[0].metric("Students", len(df))

                    cols[1].metric("Male", int((genders == "male").sum()))

                    cols[2].metric("Female", int((genders == "female").sum()))

                    cols[3].metric("Completed", int((statuses == "completed").sum()))

                    cols[4].metric("Dropouts", int((statuses == "dropped").sum()))

                    st.info("Reporting covers all years. Completed learners should have a completion month (April, June, August or December); dropped learners should have a dropout quarter (Q1–Q4). Blank status values are counted as Ongoing. Existing records are preserved; historical month/quarter values remain unspecified until updated.")

                    if st.button("Generate Excel Report", type="primary", use_container_width=True):

                              try:

                                        generated = generate_statistical_report(df, employment_df)

                                        st.session_state["excel_report_bytes"] = generated

                                        st.session_state["excel_report_filename"] = f"DBK_Statistical_Report_{datetime.now():%Y%m%d_%H%M}.xlsx"

                                        st.success("Excel report generated. Download it below.")

                              except Exception as e:

                                        st.error(f"Could not generate Excel report: {e}")

                    if st.session_state.get("excel_report_bytes"):

                              st.download_button("Download Completed Excel Report", data=st.session_state["excel_report_bytes"], file_name=st.session_state.get("excel_report_filename", "DBK_Statistical_Report.xlsx"), mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)



# =========================================================

# 16. ADMIN USER APPROVALS

# =========================================================

elif page == "User Approvals" and is_admin:

          st.subheader("Centre Account Approvals")

          with get_connection() as conn:

                    users_df = read_sql_frame(
                              conn,
                              "SELECT id,name,email,role,centre,approved,created_at FROM users ORDER BY approved ASC,name ASC",
                    )

          if users_df.empty:

                    st.info("No accounts found.")

          else:

                    display_users = users_df.copy()

                    display_users["approved"] = display_users["approved"].map({1:"Approved", 0:"Pending"})

                    st.dataframe(display_users.drop(columns=["id"]), use_container_width=True, hide_index=True)

                    pending = users_df[(users_df["approved"] == 0) & (users_df["role"] != "admin")]

                    if pending.empty:

                              st.success("There are no pending accounts.")

                    else:

                              st.write("### Pending Accounts")

                              for _, row in pending.iterrows():

                                        with st.container(border=True):

                                                  st.write(f"**{row['name']}** — {row['email']} — {row['centre']}")

                                                  approve_col, delete_col = st.columns(2)

                                                  with approve_col:

                                                            if st.button("Approve", key=f"approve_{row['id']}"):

                                                                      with get_connection() as conn:

                                                                                conn.execute("UPDATE users SET approved=1 WHERE id=?", (int(row["id"]),))

                                                                                conn.commit()

                                                                      st.success("Account approved.")

                                                                      st.rerun()

                                                  with delete_col:

                                                            if st.button("Reject / Delete", key=f"delete_{row['id']}"):

                                                                      with get_connection() as conn:

                                                                                conn.execute("DELETE FROM users WHERE id=?", (int(row["id"]),))

                                                                                conn.commit()

                                                                      st.success("Pending account deleted.")

                                                                      st.rerun()

elif page == "Change History" and is_admin:

          st.subheader("Learner and Employment Change History")
          st.caption("This audit log records who changed, imported or deleted each learner and employment record, and preserves the before-and-after data.")
          with get_connection() as conn:
                    audit_df = read_sql_frame(
                              conn,
                              """SELECT id, created_at, actor_name, actor_email, actor_role, actor_centre,
                                        action, entity_type, entity_id, entity_label, before_data, after_data, details
                                 FROM change_audit
                                 ORDER BY id DESC
                                 LIMIT 5000""",
                    )
          if audit_df.empty:
                    st.info("No learner or employment changes have been recorded yet.")
          else:
                    action_options = sorted(audit_df["action"].dropna().unique().tolist())
                    entity_options = sorted(audit_df["entity_type"].dropna().unique().tolist())
                    filter_col, entity_col, actor_col = st.columns(3)
                    with filter_col:
                              selected_actions = st.multiselect(
                                        "Filter by action",
                                        action_options,
                                        key="audit_action_filter",
                              )
                    with entity_col:
                              selected_entities = st.multiselect(
                                        "Filter by record type",
                                        entity_options,
                                        key="audit_entity_filter",
                              )
                    with actor_col:
                              actor_search = st.text_input("Search actor or record name", key="audit_actor_search")
                    filtered_audit = audit_df.copy()
                    if selected_actions:
                              filtered_audit = filtered_audit[filtered_audit["action"].isin(selected_actions)]
                    if selected_entities:
                              filtered_audit = filtered_audit[filtered_audit["entity_type"].isin(selected_entities)]
                    if actor_search.strip():
                              term = actor_search.strip().casefold()
                              actor_match = filtered_audit["actor_name"].fillna("").astype(str).str.casefold().str.contains(term, regex=False)
                              label_match = filtered_audit["entity_label"].fillna("").astype(str).str.casefold().str.contains(term, regex=False)
                              email_match = filtered_audit["actor_email"].fillna("").astype(str).str.casefold().str.contains(term, regex=False)
                              filtered_audit = filtered_audit[actor_match | label_match | email_match]
                    st.write(f"**Changes shown: {len(filtered_audit):,}** (latest 5,000 events loaded)")
                    display_columns = [
                              "created_at", "actor_name", "actor_email", "actor_role", "actor_centre",
                              "action", "entity_type", "entity_id", "entity_label",
                    ]
                    st.dataframe(
                              filtered_audit[display_columns],
                              use_container_width=True,
                              hide_index=True,
                    )
                    st.download_button(
                              "Download Filtered Change History (CSV)",
                              filtered_audit.to_csv(index=False).encode("utf-8"),
                              "DBK_Change_History.csv",
                              "text/csv",
                              key="audit_history_download",
                    )
                    detail_ids = filtered_audit["id"].astype(int).tolist()
                    if detail_ids:
                              selected_audit_id = st.selectbox(
                                        "Inspect a change",
                                        options=detail_ids,
                                        format_func=lambda audit_id: (
                                                  f"#{audit_id} · "
                                                  f"{filtered_audit.loc[filtered_audit['id'] == audit_id, 'action'].iloc[0].title()} · "
                                                  f"{filtered_audit.loc[filtered_audit['id'] == audit_id, 'entity_label'].iloc[0]}"
                                        ),
                                        key="audit_detail_id",
                              )
                              detail = filtered_audit[filtered_audit["id"] == selected_audit_id].iloc[0]
                              for key, title in (
                                        ("before_data", "Before"),
                                        ("after_data", "After"),
                                        ("details", "Import / operation details"),
                              ):
                                        raw_value = detail[key]
                                        if pd.notna(raw_value) and raw_value:
                                                  with st.expander(title):
                                                            st.json(json.loads(raw_value))

elif page == "Delete All Data" and is_admin:

          st.subheader("Delete All Learner and Employment Data")
          st.warning(
                    "This permanently deletes every learner record and every employment outcome record. "
                    "User accounts will be kept. This action cannot be undone."
          )
          count_left, count_right = st.columns(2)
          count_left.metric("Learner records to delete", len(df_all))
          count_right.metric("Employment records to delete", len(employment_all))

          confirmation_phrase = "DELETE ALL LEARNER AND EMPLOYMENT DATA"
          with st.form("delete_all_data_form"):
                    typed_confirmation = st.text_input(
                              f"Type {confirmation_phrase} to confirm",
                              help="This action affects all centres and all years.",
                    )
                    confirm_all_data_delete = st.checkbox(
                              "I understand this permanently removes all learner and employment records."
                    )
                    delete_all_data = st.form_submit_button(
                              "Permanently Delete All Learner and Employment Data",
                              type="primary",
                              use_container_width=True,
                    )

          if delete_all_data:
                    if typed_confirmation != confirmation_phrase or not confirm_all_data_delete:
                              st.error("Enter the exact confirmation phrase and check the confirmation box.")
                    else:
                              try:
                                        with get_connection() as conn:
                                                  conn.execute("BEGIN IMMEDIATE")
                                                  learner_count = conn.execute("SELECT COUNT(*) FROM learners").fetchone()[0]
                                                  employment_count = conn.execute("SELECT COUNT(*) FROM employment").fetchone()[0]
                                                  learner_records = conn.execute("SELECT * FROM learners").fetchall()
                                                  employment_records = conn.execute("SELECT * FROM employment").fetchall()
                                                  for learner_record in learner_records:
                                                            record_audit(
                                                                      conn,
                                                                      user,
                                                                      "delete",
                                                                      "learner",
                                                                      learner_record["id"],
                                                                      learner_record["name"],
                                                                      before_data=audit_record_values(learner_record),
                                                                      details={"reason": "administrator deleted all data"},
                                                            )
                                                  for employment_record in employment_records:
                                                            record_audit(
                                                                      conn,
                                                                      user,
                                                                      "delete",
                                                                      "employment",
                                                                      employment_record["id"],
                                                                      employment_record["graduate_name"],
                                                                      before_data=audit_record_values(employment_record),
                                                                      details={"reason": "administrator deleted all data"},
                                                            )
                                                  deleted_learners = conn.execute("DELETE FROM learners").rowcount
                                                  deleted_employment = conn.execute("DELETE FROM employment").rowcount
                                                  if deleted_learners != learner_count or deleted_employment != employment_count:
                                                            raise sqlite3.DatabaseError(
                                                                      "Deleted row counts did not match the record counts; the transaction was rolled back."
                                                            )
                                                  conn.commit()
                                        st.success(
                                                  f"Deleted {deleted_learners} learner record(s) and "
                                                  f"{deleted_employment} employment outcome record(s). User accounts were kept."
                                        )
                                        st.rerun()
                              except DB_ERRORS as error:
                                        st.error(f"Could not delete all learner and employment data: {error}")



# =========================================================

# 17. SIDEBAR FOOTER

# =========================================================

st.sidebar.markdown(
          '<div class="meal-nav-footer">Don Bosco Kakuma<br>Monitoring · Evaluation · Accountability · Learning</div>',
          unsafe_allow_html=True,
)
